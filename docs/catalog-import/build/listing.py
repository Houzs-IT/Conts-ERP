"""SKU listing for the owner from build/listing.mjs output.

    python -I listing.py <in.json> <out.xlsx> <source label> [<what it is after>]
"""
import sys, json, openpyxl
from openpyxl.styles import Font, PatternFill
rows = json.load(open(sys.argv[1], encoding='utf-8'))
after = sys.argv[4] if len(sys.argv) > 4 else 'migration 0132 (catalog fix of 21 Oct 2026)'
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'SKUs'
head = ['SKU code', 'Name', 'Category', 'Sub-category', 'Unit', 'Unit (meaning)', 'Model', 'Size', 'Status', 'Main supplier',
        'Supplier item code', 'Cost (RM)', 'Price 1 cost (RM)', 'Cost by seat height / tier', 'Price status']
ws.append(head)
for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDE6F0')

def matrix_text(m):
    if not isinstance(m, dict):
        return None
    cells = [(h, v) for h, v in m.items() if isinstance(v, dict)]
    if not cells:
        return None
    return '; '.join(f'{h}": ' + ' / '.join(f'{t} {v[t] / 100:,.2f}' for t in sorted(v)) for h, v in cells)

for r in rows:
    cost = (r['unit_price_sen'] if r['unit_price_sen'] is not None else r['cost_price_sen']) or 0
    note = r.get('link_note') or ''
    if r['supplier'] == 'Customer':
        status = 'customer-supplied (cost 0 by design)'
    elif cost > 0:
        status = 'cost from supplier document'
    else:
        status = 'PRICE NEEDED' if r['supplier'] else 'no supplier (generic SKU)'
    p1 = (r['price_matrix'] or {}).get('P1') if isinstance(r.get('price_matrix'), dict) else None
    ws.append([r['code'], r['name'], r['category'], r['sub_category'], r['uom'], r['unit'], r['base_model'], r['size_label'],
               r['status'].title(), r['supplier'], r['supplier_sku'], round(cost / 100, 2) if cost else None,
               round(p1 / 100, 2) if p1 else None, matrix_text(r.get('price_matrix')), status])
for col, w in zip('ABCDEFGHIJKLMNO', [30, 52, 14, 22, 7, 22, 16, 12, 9, 32, 22, 10, 10, 60, 30]): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2, min_col=12, max_col=13):
    for c in row: c.number_format = '#,##0.00'
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
rm = wb.create_sheet('Read me', 0)
for line in [f'Every SKU in the engine catalog after {after}: {len(rows)} SKUs.',
             f'Source: {sys.argv[3]}.',
             'Main supplier and Cost come from the SKU\'s main supplier link (Products > SKU Master shows the same); picking the SKU on a '
             'purchase order fills both.',
             'Cost (RM): the flat cost (bedframe: Price 2; sofa: P2 at the smallest quoted seat height). Price 1 cost: the cheaper fabric tier. '
             'Cost by seat height / tier: the sofa matrix the purchase order and sales order read when a seat height is picked.',
             'PRICE NEEDED: the supplier document has no price (DHI, LB catalogues, 2 K & K items, 39 AKEMI mattresses). '
             'Customer-supplied: the CUS- items, cost 0 on purpose.']:
    rm.append([line])
rm.column_dimensions['A'].width = 140
wb.save(sys.argv[2]); print(len(rows), 'rows')
