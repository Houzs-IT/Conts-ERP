"""rows.json -> catalog-review.xlsx, catalog-import.csv, catalog-import-excluded.csv, catalog-data.json"""
import json, sys, csv, re
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.stdout.reconfigure(encoding='utf-8')
d = json.load(open(sys.argv[1], encoding='utf-8'))
OUT = sys.argv[2]
rows, SUP, NEW_SUBS, NEW_UOMS = d['rows'], d['suppliers'], d['new_subs'], d['new_uoms']
GROUP_LABEL = {'WORKS': 'Works', 'BUILTIN': 'Built-in', 'FURNITURE': 'Furniture', 'LIGHTING': 'Lighting', 'APPLIANCE': 'Appliances', 'SOFT': 'Soft furnishing', 'SERVICE': 'Service'}
EXISTING_SUBS = {  # seeded by migrations 0061 + 0085 (production may have more added in the app)
 'WORKS': ['PROTECTION', 'HACKING', 'PLUMBING', 'WATERPROOFING', 'ELECTRICAL', 'MASONRY_TILING', 'PLASTER_CEILING', 'PARTITION', 'ALUMINIUM_GLASS', 'DOOR', 'GRILL',
           'PAINTING', 'TEXTURED_PAINTING', 'FEATURE_WALL', 'FLOORING', 'AIRCON_INSTALL', 'CLEANING', 'PROJECT_MGMT'],
 'BUILTIN': ['KITCHEN_CABINET', 'WARDROBE', 'TV_CABINET', 'SHOE_CABINET', 'STUDY_TABLE', 'COUNTERTOP', 'SHELVING'],
 'FURNITURE': ['SOFA', 'BEDFRAME', 'MATTRESS', 'DINING', 'TABLES', 'LOOSE_STORAGE', 'ACCESSORY'],
 'LIGHTING': ['DOWNLIGHT', 'PENDANT', 'LED_STRIP', 'CEILING_FAN'],
 'APPLIANCE': ['TV', 'FRIDGE', 'WASHER', 'WATER_HEATER', 'HOOD_HOB', 'AIRCON_UNIT'],
 'SOFT': ['CURTAIN_BLINDS', 'BED_LINEN', 'AMENITIES', 'DECORATION', 'CARPET', 'DIFFUSER'],
 'SERVICE': ['DELIVERY', 'STYLING', 'PHOTOSHOOT', 'AIRBNB_SETUP'],
}
EXISTING_UOMS = {'PCS', 'UNIT', 'SET', 'LOT', 'FT', 'SQFT', 'POINT', 'FULL'}
new_sub_codes = {(g, c) for g, c, _ in NEW_SUBS}
new_uom_codes = {u[0] for u in NEW_UOMS}
FOLDER = {'UNILOC': 'WORK/FLOORING/UNILOC + WORK/FEATURE WALL', 'DHI': 'WORK/FLOORING/DHI', 'LB': 'BUILT IN/STONE', 'HOOKKA': "FURNITURE/HOOKA '",
          'LOTWIN': 'LIGHTING/LOT WIN', 'WORLDONE': 'LIGHTING/WORLD ONE', 'KK': 'WORK/PLASTER CEILING + WORK (quotation QT26-13940)'}

# cross-supplier look-alikes worth a look (same product sold by two suppliers)
DUP_GROUPS = [
 ('FLR-UNI-VT230', 'GLS-KNK-031', 'V-Tech VT-230 adhesive/silicone: UNILOC RM7.50 vs K&K RM8.60'),
]
dupnote = defaultdict(list)
for a, b, why in DUP_GROUPS:
    dupnote[a].append(f"Possible same item as {b} ({why})"); dupnote[b].append(f"Possible same item as {a} ({why})")
# same colour name used by two LB codes
lbnames = defaultdict(list)
for r in rows:
    if r['supplier'] == 'LB':
        m = re.match(r'LB .*? (?:KPQ|PS|KS|TS|UV|T|M|SS)\S+ (.*?)( \(bookmatch\))?$', r['name'])
        if m and m.group(1): lbnames[m.group(1).upper()].append(r)
for nm, rs in lbnames.items():
    codes = sorted({x['supplier_item_code'] for x in rs})
    if len(codes) > 1:
        for x in rs: dupnote[x['sku']].append(f"Colour name '{nm.title()}' is also used by LB {', '.join(c for c in codes if c != x['supplier_item_code'])}")

def csv_status(r):
    if not r['importable']: return False, r['exclude_reason']
    if r['uom'] not in EXISTING_UOMS: return False, f"unit {r['uom']} is not in the engine's unit list yet (the import script adds it; the screen import cannot)"
    if (r['category'], r['sub_category']) in new_sub_codes: return False, f"sub-category {r['sub_category']} does not exist yet (the import script adds it)"
    return True, ''

for r in rows:
    r['dup'] = '; '.join(dupnote.get(r['sku'], []))
    ok, why = csv_status(r)
    r['in_csv'], r['csv_reason'] = ok, why

wb = Workbook()
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='305496')
def sheet(ws, header, data, widths):
    ws.append(header)
    for c in ws[1]: c.font, c.fill, c.alignment = H, HF, Alignment(wrap_text=True, vertical='top')
    for row in data: ws.append(row)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

# ---- Read me
ws = wb.active; ws.title = 'Read me'
readme = [
 ['Conts renovation catalog - first import review (built 20 Oct 2026 from the 18 supplier PDFs in the owner\'s IT folder)'],
 [''],
 ['What this is', 'Every item found in the supplier documents, one row per item/variant, mapped to the engine\'s category tree (Works, Built-in, Furniture, Lighting, Appliances, Soft furnishing, Service). Nothing here is invented: each row names the file and page it came from.'],
 ['SKU code scheme', '<TRADE>-<SUPPLIER>-<supplier code or model>.  TRADE: FLR flooring, FWL feature wall, CPT carpet, PTN partition, PLC plaster ceiling, DOR door, GLS glass/aluminium, PNT painting, MSN masonry, CTP countertop stone, SOF sofa, BED bedframe, ACC accessory, LED LED strip/profile, DLT downlight, TRK track light, FAN ceiling fan, SVC service, TLS tool/consumable, PLY plywood.'],
 ['', 'SUPPLIER: UNI = UNILOC (Crafted Floors Sdn Bhd), DHI = DHI Floor, LB = Little Builder, HKA = Hookka Industries, LTW = Lot Win LED, W1L = World One Lighting, KNK = K & K Gypsum.'],
 ['', 'Last part: the supplier\'s own code with spaces removed (SD40119, HSW01, KPQ001, 5531, CF/F5DC/56/PRO/MB -> F5DC-56-MB). Where the supplier gives no code (K&K price list, UNILOC price-list-only lines) a short readable name or a running number in list order is used (PLC-KNK-024). The supplier\'s exact code is always kept in "Supplier item code".'],
 ['', 'LB bookmatch slabs get -BM (CTP-LB-M9006-BM).'],
 ['Cost price', 'RM per the SKU\'s unit, from the supplier document named in "Price source". Price type says what kind of document it is: dealer price list (UNILOC), supplier price list (K&K, pages dated May-Oct 2026), supplier invoice (Lot Win), supplier sales order (World One). The K&K quotation QT26-13940 matched 10 K&K price-list lines at the same prices; its item codes are kept on those rows.'],
 ['Selling price', 'None of the documents gives a selling price, so the List/selling price column is empty everywhere. The "2990" sofa sheet may mean RM2,990 per sofa - see Questions.'],
 ['Price needed', 'DHI, LB stone and HOOKKA documents are catalogues without prices: those rows have no cost and say "Price needed".'],
 ['Import?', 'Y = goes into the system. N = left out on purpose (contractor tools, screws/nails, services, deposits, plywood) - see the reason column; the owner can switch any to Y.'],
 ['In screen CSV?', 'catalog-import.csv is for the SKU Master > Import button. That button can only use units and sub-categories that already exist and cannot set suppliers or cost, so rows needing the new units BOX/BAG/ROLL or the new sub-category Track light are left out of it (listed in catalog-import-excluded.csv). The import script (scripts/catalog-first-import.mjs) does everything: units, sub-category, suppliers, SKUs, supplier links and cost.'],
 ['Duplicates with what is already in the system', 'Column "Exists in system?" is empty: fill it from the SQL in docs/catalog-import/README.md (run read-only against production) - the script also prints every code that already exists and skips it.'],
 ['Photos', 'The engine keeps Built-in, Furniture, Lighting and Soft furnishing SKUs Inactive until a photo is added (a database rule). Works SKUs are created Active.'],
]
for r in readme: ws.append(r)
ws.column_dimensions['A'].width = 26; ws.column_dimensions['B'].width = 140
for row in ws.iter_rows():
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
ws['A1'].font = Font(bold=True, size=13)

# ---- Categories
cnt = Counter((r['category'], r['sub_category']) for r in rows)
sups = defaultdict(set)
for r in rows: sups[(r['category'], r['sub_category'])].add(FOLDER[r['supplier']])
data = []
for g in ['WORKS', 'BUILTIN', 'FURNITURE', 'LIGHTING', 'APPLIANCE', 'SOFT', 'SERVICE']:
    subs = list(EXISTING_SUBS[g]) + [c for gg, c, _ in NEW_SUBS if gg == g]
    for s in subs:
        if (g, s) not in cnt and (g, s) not in new_sub_codes: continue
        imp = sum(1 for r in rows if r['category'] == g and r['sub_category'] == s and r['importable'])
        data.append([GROUP_LABEL[g], g, s.replace('_', ' ').title() if (g, s) not in new_sub_codes else dict(((a, b), c) for a, b, c in NEW_SUBS)[(g, s)], s,
                     'NEW - import script adds it' if (g, s) in new_sub_codes else 'exists (seeded by migration 0061/0085)',
                     cnt[(g, s)], imp, '; '.join(sorted(sups[(g, s)]))])
sheet(wb.create_sheet('Categories'), ['Category', 'Category code', 'Sub-category', 'Sub-category code', 'New or existing', 'SKU rows', 'Rows to import', 'Supplier folders feeding it'],
      data, [16, 14, 26, 20, 34, 10, 12, 60])
ws = wb['Categories']
ws.append([]); ws.append(['Units (UOM)', '', 'Code', '', 'New or existing', 'SKU rows'])
for u in sorted({r['uom'] for r in rows}):
    ws.append(['', '', u, '', 'NEW - import script adds it' if u in new_uom_codes else 'exists (migration 0118)', sum(1 for r in rows if r['uom'] == u)])

# ---- SKUs
hdr = ['Category', 'Sub-category', 'Supplier', 'Supplier item code', 'Proposed SKU code', 'Name', 'Spec (size / thickness / colour / finish / packing)', 'UOM',
       'Unit on document', 'Cost price (RM)', 'List / selling price (RM)', 'Price source (file + page)', 'Price type', 'Price date', 'Item source (file + page)',
       'Confidence', 'Notes', 'Possible duplicate', 'Import? (Y/N)', 'Why not imported', 'In screen CSV?', 'Exists in system? (fill from SQL)']
data = [[GROUP_LABEL[r['category']], r['sub_category'], SUP[r['supplier']]['name'], r['supplier_item_code'], r['sku'], r['name'], r['spec'], r['uom'], r['doc_unit'],
         r['cost'], r['sell'], r['price_src'], r['price_type'], r['price_date'], r['item_src'], r['confidence'], r['notes'], r['dup'],
         'Y' if r['importable'] else 'N', r['exclude_reason'], 'Y' if r['in_csv'] else 'N', ''] for r in rows]
sheet(wb.create_sheet('SKUs'), hdr, data, [12, 16, 30, 18, 24, 44, 50, 7, 14, 10, 10, 50, 22, 11, 50, 10, 60, 40, 8, 30, 8, 14])
ws = wb['SKUs']
for c in ws['J'][1:]: c.number_format = '#,##0.00'
yellow = PatternFill('solid', fgColor='FFF2CC'); grey = PatternFill('solid', fgColor='EDEDED')
for i, r in enumerate(rows, start=2):
    if r['cost'] is None: ws.cell(i, 10).fill = yellow
    if not r['importable']:
        for c in ws[i]: c.fill = grey

# ---- Suppliers
data = []
for k, s in SUP.items():
    n = sum(1 for r in rows if r['supplier'] == k)
    data.append([s['code'], s['name'], s.get('registration_no', ''), s.get('tin_number', ''), ', '.join(x for x in [s.get('address1', ''), s.get('address2', '')] if x),
                 s.get('phone', ''), s.get('email', ''), s['note'], s['src'], n])
sheet(wb.create_sheet('Suppliers'), ['Supplier code (proposed)', 'Name', 'Registration no', 'TIN', 'Address', 'Phone', 'Email', 'Note', 'Taken from', 'SKU rows'], data,
      [18, 44, 24, 16, 60, 18, 34, 60, 50, 9])

# ---- Questions
n = lambda f: sum(1 for r in rows if f(r))
Q = [
 ('Units', f"Add three units to SKU Master > Units: BOX ({n(lambda r: r['uom']=='BOX')} SKU), BAG ({n(lambda r: r['uom']=='BAG')}) and ROLL ({n(lambda r: r['uom']=='ROLL')})? Cement/sand/rockwool/adhesive are sold by the bag, underlay/gasket by the roll. The import script adds them unless you say no.", 'K&K list, UNILOC list p3'),
 ('Sub-category', f"Add Lighting > Track light for the {n(lambda r: r['sub_category']=='TRACK_LIGHT')} Lot Win track / track light / power track items? (Alternative: put them under Downlight.)", 'Lot Win INV 0231026'),
 ('Selling prices', 'No document gives a selling price. All SKUs are created with cost only. Do you want a standard mark-up applied, or will you set selling prices by hand?', 'all'),
 ('Price needed', f"{n(lambda r: r['cost'] is None and r['importable'])} SKUs have no cost: DHI flooring + accessories (116), LB stone slabs (223), HOOKKA furniture (55) and 2 K&K glass lines. Can you get the DHI, LB and HOOKKA price lists?", 'DHI, LB, HOOKKA catalogues'),
 ('Stone unit', 'LB stone is created per slab (PCS). Do you buy per slab but sell per foot run of countertop? If so the selling SKU should be a separate Built-in > Countertop item priced per FT.', 'LB catalogues'),
 ('HOOKKA sizes', 'The HOOKKA catalogue says how many sizes each model has (e.g. 1003: 14 SKUs) but the list is cut off in the PDF. One SKU per model was created; sizes go on the order. Do you want one SKU per size instead? Then we need HOOKKA\'s full size list.', 'HOOKA.pdf p1-6'),
 ('2990 sofa sheet', "Does '2990 ALL SOFA' mean every sofa sells at RM2,990? 12 models on it (9036, 9053, 9070, 9071, 8020, 5119, 8019, 5080, 8027, 9050, 9038, 5095) are NOT in the HOOKKA catalogue - are they HOOKKA's? And what is 'Akemi Uborr' (no code)?", '2990_sofa list.pdf p1-3'),
 ('Herringbone', 'UNILOC herringbone: catalogue says 610 x 125 x 4.5mm, 16pcs = 13.2 sqft; price list says 5mm x 128 x 640mm, 16pcs = 14.1 sqft at RM3.70/sqft. Same product?', 'UNILOC price list p1, Herringbone PDF'),
 ('Fluted panel', 'Small Wave panel: price list says 12 x 204mm, catalogue 12 x 197.4mm. And is the stainless steel joint RM40.00 per piece or per box of 100?', 'UNILOC price list p2'),
 ('Stair nosing', 'UNILOC ORI SPC stair nosing 5mm/6mm: RM39 / RM42 is in the RM/sqft column but has no carton price - assumed per piece. Correct?', 'UNILOC price list p1'),
 ('Ceiling fans', 'World One charged RM130.00 for both the 56" and the 46" DEKA F5DC fan. Is that the normal price for both sizes?', 'SO_W037120, SO_W037186'),
 ('Same item, two suppliers', 'V-Tech VT-230 adhesive is on both lists: UNILOC RM7.50 and K&K RM8.60. Which supplier is main? (Created as two SKUs now.)', 'UNILOC p2, K&K p2'),
 ('Same name, two codes (LB)', 'LB uses some colour names twice: KS021 and KS023 are both "3D Carara"; Colton (T0019 / SS011); Diamond Meteor (SS004 / SS013); Calacatta White (T0001 / M9006); Pandora (M9018 sintered / TS002 translucent). Kept as separate SKUs - see "Possible duplicate".', 'LB catalogues'),
 ('Not imported: tools', f"{n(lambda r: r['exclude_reason'].startswith('Contractor tool'))} K&K lines are contractor tools or screws/nails/wall plugs (hammers, drills, lasers, brushes, gloves). Left out of the catalog - OK? (Fiber tape and screw box are kept because they were on the K&K quotation.)", 'K&K list p2-p3'),
 ('Not imported: services', 'K&K transport, unloading, night shift and scaffolding/timber/chair lend deposits (13 lines) are charges, not products. Create them as Service SKUs, or leave them for the purchase invoice?', 'K&K list p2'),
 ('Plywood', 'K&K plywood 3 / 5.2 / 9 / 12mm (4 lines) - no matching sub-category (it is carpentry board). Leave out, or add a sub-category?', 'K&K list p4'),
 ('Install materials', 'UNILOC glue, MAPEI levelling/primer, underlay, protection mats, Donson carpet glue (31 lines) are installation materials. Keep them in the catalog (so they can be bought on a PO), or leave out?', 'UNILOC list p2-p3'),
 ('Carpet', 'UNILOC carpet tiles and grass carpet (11 lines) are filed under Soft furnishing > Carpet. Or Works > Flooring?', 'UNILOC list p2'),
 ('K&K blocks', 'K&K lists ceiling boards (Unitone, PVC ceiling) inside its Partition block. They are filed under Partition as K&K groups them - move to Plaster ceiling?', 'K&K list p1'),
 ('K&K page dates', 'K&K list pages are dated 11 Sep (p1), 15 Jun (p2), 1 May (p3) and 13 Oct 2026 (p4). Prices from p2/p3 may be out of date.', 'K&K list'),
 ('Scans', 'Both Lot Win invoices (CamScanner) were readable. Two codes were read by eye from the scan: IMIDL RS433/20W LED 4K and IMI TR16/12W WH 4K - please confirm the letters.', 'Lot Win scans'),
 ('No colour names', 'DHI Sandstone 8563/8564/8565/8567 and LB UV008/UV009/UV071 have no colour names in the catalogue.', 'DHI p12, LB quartz p29'),
 ('Photos', f"{n(lambda r: r['importable'] and r['category'] in ('BUILTIN','FURNITURE','LIGHTING','SOFT'))} Built-in/Furniture/Lighting/Soft SKUs will stay Inactive until a photo is added (engine rule). The catalogues have the photos - should we cut them out and attach them?", 'engine rule (migration 0062)'),
]
sheet(wb.create_sheet('Questions'), ['#', 'Topic', 'Question for the owner', 'Where'], [[i + 1, t, q, w] for i, (t, q, w) in enumerate(Q)], [4, 22, 120, 34])
for row in wb['Questions'].iter_rows():
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')

# ---- Summary
data = []
for k in SUP:
    rs = [r for r in rows if r['supplier'] == k]
    data.append([SUP[k]['name'], len(rs), sum(r['importable'] for r in rs), sum(r['cost'] is not None for r in rs), sum(r['cost'] is None for r in rs if r['importable']),
                 sum(not r['importable'] for r in rs), sum(r['in_csv'] for r in rs)])
data.append(['TOTAL', len(rows), sum(r['importable'] for r in rows), sum(r['cost'] is not None for r in rows), sum(r['cost'] is None for r in rows if r['importable']),
             sum(not r['importable'] for r in rows), sum(r['in_csv'] for r in rows)])
data.append([]); data.append(['By category', 'SKU rows', 'To import', 'With cost', 'Price needed (to import)', 'Not imported', 'In screen CSV'])
for g in GROUP_LABEL:
    rs = [r for r in rows if r['category'] == g]
    if rs: data.append([GROUP_LABEL[g], len(rs), sum(r['importable'] for r in rs), sum(r['cost'] is not None for r in rs), sum(r['cost'] is None for r in rs if r['importable']),
                        sum(not r['importable'] for r in rs), sum(r['in_csv'] for r in rs)])
data.append([]); data.append(['By confidence', 'SKU rows'])
for c in ['high', 'medium', 'low']: data.append([c, sum(r['confidence'] == c for r in rows)])
sheet(wb.create_sheet('Summary'), ['Supplier', 'SKU rows', 'To import', 'With cost', 'Price needed (to import)', 'Not imported', 'In screen CSV'], data, [44, 10, 10, 10, 22, 13, 13])
wb.move_sheet('Summary', offset=-4)
wb.save(f"{OUT}/catalog-review.xlsx")

# ---- screen CSV (exact importer headings) + excluded list
with open(f"{OUT}/catalog-import.csv", 'w', newline='', encoding='utf-8-sig') as fh:
    w = csv.writer(fh)
    w.writerow(['code', 'name', 'category', 'sub_category', 'uom', 'status', 'description', 'branding'])
    for r in rows:
        if r['in_csv']:
            w.writerow([r['sku'], r['name'], r['category'], r['sub_category'], r['uom'], 'Active' if r['category'] in ('WORKS', 'SERVICE') else 'Inactive',
                        r['spec'][:500], SUP[r['supplier']]['name']])
with open(f"{OUT}/catalog-import-excluded.csv", 'w', newline='', encoding='utf-8-sig') as fh:
    w = csv.writer(fh)
    w.writerow(['code', 'name', 'reason'])
    for r in rows:
        if not r['in_csv']: w.writerow([r['sku'], r['name'], r['csv_reason']])

# ---- data file for the import script
data = dict(built='2026-10-20', source='owner IT folder, 18 PDFs', suppliers=SUP, new_sub_categories=[dict(group=g, code=c, label=l) for g, c, l in NEW_SUBS],
            new_uoms=[dict(code=c, label=l, print_label=p, measure=m) for c, l, p, m in NEW_UOMS],
            skus=[dict(code=r['sku'], name=r['name'], category=r['category'], sub_category=r['sub_category'], uom=r['uom'], description=r['spec'],
                       supplier=r['supplier'], supplier_item_code=r['supplier_item_code'], cost_rm=r['cost'], price_date=r['price_date'] or None,
                       price_source=r['price_src'], price_type=r['price_type'], item_source=r['item_src'], confidence=r['confidence'], notes=r['notes'])
                  for r in rows if r['importable']])
json.dump(data, open(f"{OUT}/catalog-data.json", 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print('rows', len(rows), 'csv', sum(r['in_csv'] for r in rows), 'import', len(data['skus']))
