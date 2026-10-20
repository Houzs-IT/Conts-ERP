"""Catalog fix of 21 Oct 2026 (owner feedback on the first import): bedframes split by size, sofas split by compartment,
the 13 "2990 list" sofas removed, every SKU's unit checked against its supplier document.

    python -I docs/catalog-import/build/fix.py <repo root>

Reads docs/catalog-import/catalog-review.xlsx (sheet SKUs, the rows imported) and catalog-data.json. Writes
docs/catalog-import/catalog-fix-data.json, docs/catalog-import/catalog-fix-review.xlsx and
migrations-pg/0132_catalog_fix_sizes_compartments_units.sql. Edit this script, not the migration.

Unit rule (owner via lead, 21 Oct 2026): the SKU takes the unit the supplier writes on its own document, mapped to the
matching scm.product_uoms code; a code is added only for a unit a supplier writes that the list lacks (K & K's "no" /
"nos" -> NOS). Where the document writes no unit at all the SKU keeps its unit and goes on the "Owner to decide" sheet.
"""
import sys, os, json, re, hashlib
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

ROOT = sys.argv[1]
DOCS = os.path.join(ROOT, 'docs', 'catalog-import')
MIG = os.path.join(ROOT, 'migrations-pg', '0132_catalog_fix_sizes_compartments_units.sql')
TAG = 'catalog-fix 2026-10'
HKA_FILE = "FURNITURE/HOOKA '/HOOKA.pdf"
HKA_2990 = "FURNITURE/HOOKA '/2990_sofa list.pdf"

wb = openpyxl.load_workbook(os.path.join(DOCS, 'catalog-review.xlsx'), read_only=True)
rows = list(wb['SKUs'].iter_rows(values_only=True))
H = rows[0]
REV = {r[H.index('Proposed SKU code')]: dict(zip(H, r)) for r in rows[1:] if r[H.index('Import? (Y/N)')] == 'Y'}
DATA = json.load(open(os.path.join(DOCS, 'catalog-data.json'), encoding='utf-8'))
SKUS = DATA['skus']
assert len(SKUS) == 819 and all(k['code'] in REV for k in SKUS)

# ---------------------------------------------------------------- units: labels and the one added code
UOM_LABELS = [  # code, label, print label (only codes that exist; owner-added codes are left alone)
    ('PCS', 'Piece', 'PCS'), ('UNIT', 'Unit', 'UNIT'), ('SET', 'Set', 'SET'), ('LOT', 'Lump sum (lot)', 'LS'),
    ('FT', 'Foot run (linear foot)', 'FT RUN'), ('SQFT', 'Square foot', 'SQ FT'), ('POINT', 'Point', 'POINT'),
    ('FULL', 'Full unit', 'FULL UNIT'), ('BOX', 'Box / carton', 'BOX'), ('BAG', 'Bag', 'BAG'), ('ROLL', 'Roll', 'ROLL'),
]
OLD_LABELS = {'PCS': ('Pieces', 'PCS'), 'UNIT': ('Unit', 'UNIT'), 'SET': ('Set', 'SET'), 'LOT': ('Lump sum', 'LS'),
              'FT': ('Feet (run)', 'FT'), 'SQFT': ('Square feet', 'SQFT'), 'POINT': ('Point', 'POINT'),
              'FULL': ('Full unit', 'FULL UNIT'), 'BOX': ('Box / carton', 'BOX'), 'BAG': ('Bag', 'BAG'), 'ROLL': ('Roll', 'ROLL')}
NEW_UOMS = [('NOS', 'Nos (number)', 'NOS', False)]  # K & K writes "no" / "nos" (doors, silicone, door closer)

# ---------------------------------------------------------------- unit decision per SKU
def short(src):
    return src.replace('WORK/FLOORING/UNILOC/', '').replace('WORK/FEATURE WALL/', '').replace('WORK/PLASTER CEILING/', '') \
              .replace('WORK/FLOORING/DHI/', '').replace('BUILT IN/STONE/', '').replace('FURNITURE/HOOKA \'/', '') \
              .replace('LIGHTING/LOT WIN/', '').replace('LIGHTING/WORLD ONE/', '')

KK_QT = {  # K & K's own quotation QT26-13940 (12 Aug 2026) prints a unit for these list lines that print none
    'PLC-KNK-133': ('PCS', 'QT26-13940 line 9 "K- FIBER TAPE 40MMX50M (1.5\'\') ... 1.00 PCS"'),
    'PNT-KNK-001': ('UNIT', 'QT26-13940 line 6 "G- MAXILITE (WHITE 15245) 18 LITER 1.00 UNIT"'),
}
KK_MAP = {'pcs': 'PCS', 'PCS': 'PCS', 'set': 'SET', 'bag': 'BAG', 'roll': 'ROLL', 'ft': 'FT', 'sq ft': 'SQFT', 'no': 'NOS', 'nos': 'NOS'}

def decide(k):
    """-> (new unit, evidence, confidence, owner_question_key or None)"""
    r = REV[k['code']]
    du = r['Unit on document']
    src = r['Price source (file + page)'] or r['Item source (file + page)']
    sup = k['supplier']
    old = k['uom']
    if sup == 'UNILOC':
        if du == 'RM/sqft (carton sold)':
            return 'SQFT', f'0022_UNILOC Price List p1: price column "RM/FT²" (also RM/CTN); catalogue {short(k["item_source"])}: "PACKING PER BOX ... SQFT/BOX"', 'high', None
        if du in ('RM/FT2', 'RM/FT2 (roll / cut)'):
            return 'SQFT', f'{short(src)}: price column "RM/FT²"', 'high', None
        if du == 'RM/FT2 column (per piece?)':
            return 'SQFT', f'{short(src)}: the price sits in the "RM/FT²" column (row says 4pcs, no RM/CTN price)', 'medium', 'nosing'
        if du in ('RM/Unit', 'RM/UNIT', 'RM/Unit (loose)'):
            col = 'RM/UNIT' if du == 'RM/UNIT' else 'RM/Unit'
            return 'UNIT', f'{short(src)}: price column "{col}"' + (' (Loose)' if 'loose' in du else ''), 'high', ('joint' if k['code'] == 'FWL-UNI-JOINT-SS' else None)
        if du == 'RM/Unit per roll (loose)':
            return 'ROLL', f'{short(src)}: columns "Sqft/Roll" and "RM/Unit Loose | Roll" (one unit is one roll)', 'high', None
    if sup == 'KK':
        if k['code'] in KK_QT:
            u, ev = KK_QT[k['code']]
            return u, f'K & K PRICE LIST.pdf prints no unit for this line; {ev}', 'high', None
        if du == 'BOX (from quotation QT26-13940)':
            return 'BOX', 'QT26-13940 line 8 "K- SCREW ... (1 CTN X 20BOX) 1.00 BOX"', 'high', None
        if du in KK_MAP:
            return KK_MAP[du], f'{short(r["Item source (file + page)"])}: unit column "{du}"', 'high', None
        if du == '(none printed)':
            if '(BAG)' in k['name']:
                return 'BAG', f'{short(r["Item source (file + page)"])}: no unit column; the item itself is written "{k["name"]}"', 'medium', None
            return old, f'{short(r["Item source (file + page)"])}: no unit printed for this line', 'none', 'kk'
    if sup == 'LOTWIN':
        return 'UNIT', f'{short(src)}: column "Price/Unit" (qty counted in units)', 'high', None
    if sup == 'WORLDONE':
        return 'UNIT', f'{short(src)}: column "U/ PriceRM" (unit price; qty counted in units)', 'high', None
    if sup == 'DHI':
        if k['uom'] == 'SQFT':
            return old, f'{short(k["item_source"])}: no price and no selling unit; only "24.2 sqft (10pcs/ctn)" packing', 'none', 'dhi_floor'
        return old, f'{short(k["item_source"])}: photo page, no unit and no price', 'none', 'dhi_acc'
    if sup == 'LB':
        return old, f'{short(k["item_source"])}: colour card with slab size only; no unit and no price', 'none', 'lb'
    if sup == 'HOOKKA':
        return old, f'{short(k["item_source"])}: lists models and "SKUs" only; no unit and no price', 'none', 'hookka'
    raise SystemExit(f'no unit rule for {k["code"]} ({sup}, {du!r})')

# ---------------------------------------------------------------- HOOKKA catalogue lines (HOOKA.pdf p1-p5, as printed)
HKA_LINES = [
    (1, '1003', '14 SKUs · 200CMX200CM, 220CMX220CM, 3.5FT, 3FT, 5FT…'),
    (1, '1005', '5 SKUs · 200CMX200CM, 3.5FT, 3FT, 5FT, 6FT'),
    (1, '1007', '12 SKUs · 152CMX200CM, 153CMX200CM, 183CMX200CM,'),
    (1, '1008', '4 SKUs · 3.5FT, 3FT, 5FT, 6FT'),
    (1, '1009', '4 SKUs · 3.5FT, 3FT, 5FT, 6FT'),
    (1, '1013', '5 SKUs · 200CMX200CM, 3.5FT, 3FT, 5FT, 6FT'),
    (1, '1019', '8 SKUs · 3.5FT, 3FT, 5FT, 6FT'),
    (1, '1023', '8 SKUs · 3.5FT, 3FT, 5FT, 6FT'),
    (1, '1030', '8 SKUs · 153x200cm, 3.5FT, 3FT, 5FT, 6FT'),
    (2, '1041', '5 SKUs · 107CMX200CM, 3.5FT, 3FT, 5FT, 6FT'),
    (2, '1051', '1 SKU · 6FT'),
    (2, '1052', '3 SKUs · 3.5FT, 5FT, 6FT'),
    (2, '2003', '5 SKUs · 183CMX200CM, 3.5FT, 3FT, 5FT, 6FT'),
    (2, '2006', '10 SKUs · 190CMX200CM, 200CMX200CM, 3.5FT, 3FT, 5FT…'),
    (2, '2008', '12 SKUs · 152CMX200CM, 183CMX200CM, 200CMX200CM,'),
    (2, '2009', '11 SKUs · 152CMX200CM, 200CMX200CM, 210CMX210CM,'),
    (2, '2010', '4 SKUs · 3.5FT, 3FT, 5FT, 6FT'),
    (2, '2011', '5 SKUs · 200CMX200CM, 3.5FT, 3FT, 5FT, 6FT'),
    (3, '2023', '4 SKUs · 3FT, 6FT'),
    (3, '2027', '4 SKUs · 3.5FT, 3FT, 5FT, 6FT'),
    (3, '2033', '8 SKUs · 3.5FT, 3FT, 5FT, 6FT'),
    (3, '2038', '7 SKUs · 3.5FT, 3FT, 5FT, 6FT'),
    (3, '2041', '4 SKUs · 3.5FT, 3FT, 5FT, 6FT'),
    (3, '2049', '1 SKU · 6FT'),
    (3, 'DIVAN', '9 SKUs · 153CMX210CM, 170CMX200CM, 200CMX200CM,'),
    (4, '5530', '13 SKUs · 1A(LHF), 1A(RHF), 1NA, 1S, 2A(LHF)…'),
    (4, '5531', '13 SKUs · 1A(LHF), 1A(RHF), 1NA, 1S, 2A(LHF)…'),
    (4, '5535', '13 SKUs · 1A(LHF), 1A(RHF), 1NA, 1S, 2A(LHF)…'),
    (4, '5536', '14 SKUs · 1A(LHF), 1A(RHF), 1NA, 1S, 2A(LHF)…'),
    (4, '5537', '13 SKUs · 1A(LHF), 1A(RHF), 1NA, 1S, 2A(LHF)…'),
    (4, '5539', '17 SKUs · 1A(LHF), 1A(RHF), 1B(LHF), 1B(RHF), 1NA…'),
    (4, '5540', '14 SKUs · 1A(LHF), 1A(RHF), 1NA, 1S, 2A(LHF)…'),
    (4, '5543', '13 SKUs · 12" X 28", 1B(LHF), 1B(RHF), 1NA, 1S…'),
    (4, '5545', '17 SKUs · 1A(LHF), 1A(RHF), 1B(LHF), 1B(RHF), 1NA…'),
    (5, '5546', '17 SKUs · 1A(LHF), 1A(RHF), 1B(LHF), 1B(RHF), 1NA…'),
    (5, '5547', '17 SKUs · 1A(LHF), 1A(RHF), 1B(LHF), 1B(RHF), 1NA…'),
    (5, '5548', '17 SKUs · 1A(LHF), 1A(RHF), 1B(LHF), 1B(RHF), 1NA…'),
    (5, '5549', '17 SKUs · 1A(LHF), 1A(RHF), 1B(LHF), 1B(RHF), 1NA…'),
    (5, '6112', '1 SKU · 2S'),
]
# The engine's bedframe size pool on production (maintenance config bedframeSizes: K, Q, S, SS, SK, SP) with the labels
# of product-models.ts SIZE_INFO. A size HOOKKA prints in cm that is not 200 x 200 is a made-to-order size: code SP.
STD = {'6FT': 'K', '5FT': 'Q', '3FT': 'S', '3.5FT': 'SS', '200CMX200CM': 'SK'}
STD_LABEL = {'K': '6FT', 'Q': '5FT', 'S': '3FT', 'SS': '3.5FT', 'SK': '200X200CM'}
SIZE_ORDER = ['K', 'Q', 'SS', 'S', 'SK', 'SP']
# The engine's sofa compartment library (scm.compartment_library, 18 rows from Houzs) and the production pool.
COMPARTMENTS = {'1A(LHF)', '1A(RHF)', '1B(LHF)', '1B(RHF)', '1NA', '2A(LHF)', '2A(RHF)', '2B(LHF)', '2B(RHF)', '2NA', 'CNR',
                'L(LHF)', 'L(RHF)', 'Console', 'STOOL', '1S', '2S', '3S'}
COMP_ORDER = ['1A(LHF)', '1A(RHF)', '1B(LHF)', '1B(RHF)', '1NA', '1S', '2A(LHF)', '2S']

def sku_id(code):
    return 'mfg-' + hashlib.sha256(f'{TAG}|{code}'.encode()).hexdigest()[:12]

by_code = {k['code']: k for k in SKUS}
new_rows, models, hka_notes = [], [], []
for page, mc, line in HKA_LINES:
    m = re.match(r'(\d+) SKUs? · (.*)$', line)
    count, rest = int(m.group(1)), m.group(2)
    cut = rest.endswith('…') or rest.endswith(',')
    names = [s.strip() for s in rest.rstrip('…').rstrip(',').split(',') if s.strip()]
    sofa = mc.startswith('55') or mc == '6112'
    old_code = f"{'SOF' if sofa else 'BED'}-HKA-{mc}"
    old = by_code[old_code]
    nick = old['name'].split(mc, 1)[1].strip()  # "Xammar" etc. (2990 list name), or ''
    seen = f'{len(names)} named of {count}' + (' (list cut off in the catalogue)' if cut else '')
    variants, skipped = [], []
    for n in names:
        if sofa:
            if n not in COMPARTMENTS:
                skipped.append(n); continue
            variants.append(dict(code=f'{old_code}-{n}', comp=n, size_code=None, size_label=None, printed=n,
                                 name=f'HOOKKA SOFA {mc}{" " + nick.upper() if nick else ""} {n}'))
        else:
            u = n.upper().replace(' ', '')
            if u in STD:
                sc = STD[u]; lbl = STD_LABEL[sc]; suffix = sc
            else:
                mm = re.match(r'^(\d+)(?:CM)?X(\d+)CM$', u)
                assert mm, n
                sc = 'SP'; lbl = f'{mm.group(1)}X{mm.group(2)}CM'; suffix = f'{mm.group(1)}X{mm.group(2)}'
            variants.append(dict(code=f'{old_code}-({suffix})', comp=None, size_code=sc, size_label=lbl, printed=n,
                                 name=f'HOOKKA BEDFRAME {mc} ({lbl})'))
    key = (lambda v: (COMP_ORDER.index(v['comp']),)) if sofa else (lambda v: (SIZE_ORDER.index(v['size_code']), v['size_label']))
    variants.sort(key=key)
    allowed = {'compartments': [v['comp'] for v in variants]} if sofa else {'sizes': sorted({v['size_code'] for v in variants}, key=SIZE_ORDER.index)}
    models.append(dict(model_code=old_code, category='SOFA' if sofa else 'BEDFRAME', allowed_options=allowed, page=page,
                       printed=line, count=count, cut=cut, created=len(variants), skipped=skipped))
    for v in variants:
        what = f'compartment {v["comp"]}' if sofa else f'size {v["printed"]}'
        new_rows.append(dict(
            id=sku_id(v['code']), code=v['code'], name=v['name'], category='SOFA' if sofa else 'BEDFRAME', base_model=old_code,
            size_code=v['size_code'], size_label=v['size_label'], old_code=old_code, page=page, printed=v['printed'],
            supplier_sku=f'{mc} {v["printed"]}',
            description=f'HOOKKA model {mc}, {what}. HOOKA.pdf p{page}: "{line}" ({seen}).',
            notes=f'{TAG}: PRICE NEEDED - no price in {HKA_FILE} p{page}; supplier model {mc}, {what}. Replaces {old_code}.'))

DELETE_2990 = ['SOF-HKA-5080', 'SOF-HKA-5095', 'SOF-HKA-5119', 'SOF-HKA-8019', 'SOF-HKA-8020', 'SOF-HKA-8027', 'SOF-HKA-9036',
               'SOF-HKA-9038', 'SOF-HKA-9050', 'SOF-HKA-9053', 'SOF-HKA-9070', 'SOF-HKA-9071', 'SOF-HKA-AKEMI-UBORR']
assert all(c in by_code for c in DELETE_2990)
REPLACED = [m['model_code'] for m in models]
assert len(REPLACED) == 39 and not set(REPLACED) & set(DELETE_2990)

# ---------------------------------------------------------------- unit audit over the SKUs that stay
audit = []
for k in SKUS:
    if k['code'] in DELETE_2990 or k['code'] in REPLACED:
        continue
    new, ev, conf, q = decide(k)
    audit.append(dict(code=k['code'], name=k['name'], supplier=DATA['suppliers'][k['supplier']]['name'], old=k['uom'], new=new,
                      evidence=ev, confidence=conf, q=q))
changes = [a for a in audit if a['new'] != a['old']]
# The new sized / compartment SKUs carry the unit of the SKU they replace (UNIT): HOOKKA prints none.
for n in new_rows:
    assert by_code[n['old_code']]['uom'] == 'UNIT'

# ---------------------------------------------------------------- outputs: data json
fix = dict(tag=TAG, units=dict(labels=UOM_LABELS, added=NEW_UOMS), unit_changes=[{k: a[k] for k in ('code', 'old', 'new')} for a in changes],
           models=models, new_skus=new_rows, replaced=REPLACED, deleted_2990=DELETE_2990)
json.dump(fix, open(os.path.join(DOCS, 'catalog-fix-data.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

# ---------------------------------------------------------------- outputs: migration
def lit(v):
    return 'null' if v is None else "'" + str(v).replace("'", "''") + "'"
def jb(v):
    s = json.dumps(v, ensure_ascii=False)
    assert '$cfx$' not in s
    return f'$cfx${s}$cfx$::jsonb'

C = "(select id from public.companies where code = 'HOUZS')"
mig_models = [dict(model_code=m['model_code'], category=m['category'], allowed_options=m['allowed_options']) for m in models]
mig_new = [{k: n[k] for k in ('id', 'code', 'name', 'category', 'base_model', 'size_code', 'size_label', 'old_code', 'supplier_sku', 'description', 'notes')} for n in new_rows]
retire = [dict(code=c, reason='2990') for c in DELETE_2990] + [dict(code=c, reason='replaced') for c in REPLACED]
n_bed = sum(1 for n in new_rows if n['category'] == 'BEDFRAME'); n_sofa = len(new_rows) - n_bed
sql = f"""-- REVERSAL: none in SQL - this deletes the 13 "2990 list" sofas and the 39 single HOOKKA bedframe / sofa SKUs (with their supplier links), which only a re-run of docs/catalog-import/sql/catalog-first-import.part-02-of-02.sql brings back; then delete the {len(new_rows)} SKUs named in docs/catalog-import/catalog-fix-data.json (new_skus) with their supplier links, set scm.product_models category back to 'FURNITURE' and allowed_options to '{{}}' for the 39 model codes in "replaced", set the units in unit_changes back to "old", delete product_uoms NOS (after no SKU uses it), and restore the unit labels listed in catalog-fix-review.xlsx (sheet Units).
-- Catalog fix after the owner's review of the first import (21 Oct 2026; docs/catalog-import/catalog-fix-review.xlsx):
--  1. HOOKKA bedframes split by size, the engine's way (product-models.ts generate-skus): the model (product_models,
--     category BEDFRAME, allowed_options.sizes) and one SKU per size, code <model>-(<size code>), base_model, size_code
--     (K / Q / SS / S / SK, SP for a made-to-order cm size) and size_label. Only the sizes HOOKA.pdf prints are created.
--  2. HOOKKA sofas split by compartment the same way: model category SOFA, allowed_options.compartments, one SKU per
--     compartment, code <model>-<compartment>. HOOKA.pdf cuts each list short ("13 SKUs · 1A(LHF), ... 2A(LHF)…"), so
--     only the compartments it prints are created; all are already in scm.compartment_library.
--     Each new SKU is linked to HOOKKA (main supplier, supplier code = model + size / compartment, cost 0, price needed).
--     The single SKU each one replaces is deleted with its supplier link when nothing else names it, else deactivated.
--  3. The 13 "2990 list" sofas are deleted with their supplier links, cost history and model, when nothing names them
--     (else deactivated). Every delete / deactivation is written to scm.entity_audit_log.
--  4. Units: clearer labels, unit NOS added (K & K writes "no" / "nos"), {len(changes)} SKUs moved to the unit their supplier
--     document writes. A SKU is changed only while it still has the unit the import gave it.
-- Built by docs/catalog-import/build/fix.py from catalog-data.json and catalog-review.xlsx; edit that script, not this file.
-- Does nothing for SKUs that do not exist (a database without the first import). Safe to run again: it finds nothing to do.
set local search_path = scm, public, extensions;

-- 4a. Unit labels (only codes that exist) and the added unit.
update scm.product_uoms u set label = v.label, print_label = v.print_label
  from (values
  {(','+chr(10)+'  ').join(f'({lit(c)}, {lit(l)}, {lit(p)})' for c, l, p in UOM_LABELS)}
  ) v(code, label, print_label)
 where u.code = v.code and (u.label, u.print_label) is distinct from (v.label, v.print_label);

insert into scm.product_uoms (code, label, print_label, measure, builtin, sort_order) values
  {(','+chr(10)+'  ').join(f"({lit(c)}, {lit(l)}, {lit(p)}, {'true' if m else 'false'}, false, 200)" for c, l, p, m in NEW_UOMS)}
on conflict (code) do nothing;

-- 4b. Each SKU to the unit its supplier writes (only while it still has the imported unit).
update scm.mfg_products p set uom = v.new, updated_at = now()
  from jsonb_to_recordset({jb(fix['unit_changes'])}) as v(code text, old text, new text)
 where p.company_id = {C} and p.code = v.code and p.uom = v.old;

-- 1-2. The single HOOKKA SKUs that are still there are the work list; nothing below runs for the others.
create temp table cfx_old on commit drop as
select p.id, p.code, p.status::text as status, p.uom, p.branding, p.model_id
  from scm.mfg_products p
 where p.company_id = {C} and p.category::text = 'FURNITURE'
   and p.code = any({lit('{' + ','.join(REPLACED) + '}')}::text[]);

create temp table cfx_new on commit drop as
select n.* from jsonb_to_recordset({jb(mig_new)}) as n(id text, code text, name text, category text, base_model text,
  size_code text, size_label text, old_code text, supplier_sku text, description text, notes text)
 where n.old_code in (select code from cfx_old);

-- The model of each single SKU (model_code = SKU code, made by the import) becomes the bedframe / sofa model.
create temp table cfx_models on commit drop as
select * from jsonb_to_recordset({jb(mig_models)}) as m(model_code text, category text, allowed_options jsonb)
 where m.model_code in (select code from cfx_old);

update scm.product_models m
   set category = x.category::scm.mfg_product_category, allowed_options = coalesce(m.allowed_options, '{{}}'::jsonb) || x.allowed_options,
       updated_at = now()
  from cfx_models x
 where m.company_id = {C} and m.model_code = x.model_code and m.category::text = 'FURNITURE'
   and not exists (select 1 from scm.product_models m2 where m2.company_id = m.company_id and m2.model_code = x.model_code
                                                        and m2.category::text = x.category);

insert into scm.product_models (company_id, model_code, name, category, branding, allowed_options)
select {C}, x.model_code, o.name, x.category::scm.mfg_product_category, o.branding, x.allowed_options
  from cfx_models x join scm.mfg_products o on o.company_id = {C} and o.code = x.model_code
on conflict (company_id, model_code, category) do nothing;

insert into scm.mfg_products (id, company_id, code, name, category, uom, status, cost_price_sen, description, branding,
                              model_id, base_model, size_code, size_label)
select n.id, {C}, n.code, n.name, n.category::scm.mfg_product_category, o.uom, o.status::scm.mfg_product_status, 0,
       n.description, o.branding, m.id, n.base_model, n.size_code, n.size_label
  from cfx_new n
  join cfx_old o on o.code = n.old_code
  join scm.product_models m on m.company_id = {C} and m.model_code = n.base_model and m.category::text = n.category
on conflict (company_id, code) do nothing;

-- Linked to the same supplier as the SKU it replaces (HOOKKA), as main supplier; no cost (the catalogue has none).
insert into scm.supplier_material_bindings
    (company_id, supplier_id, material_kind, item_code, material_name, supplier_sku, unit_price_sen, price_valid_from, is_main_supplier, notes)
select {C}, b.supplier_id, 'mfg_product', n.code, n.name, n.supplier_sku, 0, null,
       not exists (select 1 from scm.supplier_material_bindings x where x.company_id = {C} and x.material_kind = 'mfg_product'
                                                                    and x.item_code = n.code and x.is_main_supplier),
       n.notes
  from cfx_new n
  join scm.supplier_material_bindings b on b.company_id = {C} and b.material_kind = 'mfg_product' and b.item_code = n.old_code
                                       and b.is_main_supplier
 where not exists (select 1 from scm.supplier_material_bindings x where x.company_id = {C} and x.material_kind = 'mfg_product'
                                                                   and x.item_code = n.code and x.supplier_id = b.supplier_id);

-- 1-3. What goes: the 13 "2990 list" sofas and the single SKUs replaced above (only those whose new SKUs now exist).
create temp table cfx_retire on commit drop as
select p.id, p.code, r.reason, 0::bigint as refs, ''::text as used_in
  from jsonb_to_recordset({jb(retire)}) as r(code text, reason text)
  join scm.mfg_products p on p.company_id = {C} and p.code = r.code and p.category::text = 'FURNITURE'
 where (r.reason = '2990' or exists (select 1 from scm.mfg_products n where n.company_id = {C} and n.base_model = r.code))
   and not exists (select 1 from scm.entity_audit_log a where a.entity_type = 'mfg_product' and a.entity_id = p.code
                                                        and a.note like {lit(TAG + '%')});

-- Anything that names the SKU besides its own supplier link and cost history: every item-code column of both engine
-- schemas (document lines, stock, price history, quotations, pwp, racks ...), and, for a replaced SKU, a photo.
do $$
declare t record; n bigint;
begin
  for t in
    select c.table_schema as s, c.table_name as tb, c.column_name as col
      from information_schema.columns c
      join information_schema.tables x on x.table_schema = c.table_schema and x.table_name = c.table_name and x.table_type = 'BASE TABLE'
     where c.table_schema in ('scm', 'public') and c.data_type = 'text'
       and c.column_name in ('item_code', 'product_code', 'new_item_code', 'trigger_item_code', 'redeemed_item_code', 'sku', 'ac_item_code')
       and not (c.table_schema = 'scm' and c.table_name in ('supplier_material_bindings', 'supplier_binding_price_history'))
  loop
    execute format('update pg_temp.cfx_retire r set refs = r.refs + q.n, used_in = r.used_in || %L || '' ''
                      from (select x.%I as code, count(*) as n from %I.%I x where x.%I in (select code from pg_temp.cfx_retire) group by 1) q
                     where q.code = r.code', t.s || '.' || t.tb, t.col, t.s, t.tb, t.col);
  end loop;
end $$;
update cfx_retire r set refs = r.refs + q.n, used_in = r.used_in || 'scm.mfg_product_photos '
  from (select product_id, count(*) as n from scm.mfg_product_photos group by 1) q
 where r.reason = 'replaced' and q.product_id = r.id;

insert into scm.entity_audit_log (company_id, entity_type, entity_id, action, actor_name_snapshot, field_changes, source, note)
select {C}, 'mfg_product', r.code, case when r.refs = 0 then 'delete' else 'deactivate' end, 'migration 0132', '[]'::jsonb, 'migration',
       {lit(TAG + ': ')} || case when r.reason = '2990' then 'owner 21 Oct 2026: the "2990 list" sofas must not be in the system'
                                  else 'owner 21 Oct 2026: replaced by one SKU per ' ||
                                       case when r.code like 'SOF-%' then 'compartment' else 'size' end || ' (base model ' || r.code || ')' end
       || case when r.refs = 0 then '; nothing named it, deleted with its supplier link and cost history'
               else '; still named in ' || trim(r.used_in) || ', so deactivated instead of deleted' end
  from cfx_retire r;

delete from scm.supplier_binding_price_history h using cfx_retire r
 where r.refs = 0 and h.company_id = {C} and h.material_kind = 'mfg_product' and h.item_code = r.code;
delete from scm.supplier_material_bindings b using cfx_retire r
 where r.refs = 0 and b.company_id = {C} and b.material_kind = 'mfg_product' and b.item_code = r.code;
delete from scm.mfg_products p using cfx_retire r where r.refs = 0 and p.id = r.id;
update scm.mfg_products p set status = 'INACTIVE', updated_at = now() from cfx_retire r where r.refs > 0 and p.id = r.id;
-- A model left with no SKU and still a plain Furniture model existed only for a deleted SKU.
delete from scm.product_models m using cfx_retire r
 where m.company_id = {C} and m.model_code = r.code and m.category::text = 'FURNITURE'
   and not exists (select 1 from scm.mfg_products p where p.model_id = m.id);

-- No supplier link may be left pointing at a SKU this file deleted.
do $$ begin
  if exists (select 1 from scm.supplier_material_bindings b join pg_temp.cfx_retire r on r.code = b.item_code
              where r.refs = 0 and b.company_id = (select id from public.companies where code = 'HOUZS')) then
    raise exception 'catalog fix: a supplier link still names a deleted SKU';
  end if;
end $$;

insert into scm.app_config (key, value, description, company_id)
select 'note:catalog-fix-2026-10',
       json_build_object('tag', {lit(TAG)}, 'at', now(), 'unit_changes', {len(changes)}, 'units_added', array['NOS'],
                         'bedframe_skus', (select count(*) from cfx_new where category = 'BEDFRAME'),
                         'sofa_skus', (select count(*) from cfx_new where category = 'SOFA'),
                         'deleted', (select coalesce(array_agg(code order by code), '{{}}') from cfx_retire where refs = 0),
                         'deactivated', (select coalesce(array_agg(code order by code), '{{}}') from cfx_retire where refs > 0))::text,
       'Catalog fix note: {TAG} (migration 0132)', {C}
 where exists (select 1 from cfx_old) or exists (select 1 from cfx_retire)
on conflict (key) do nothing;

notify pgrst, 'reload schema';
"""
open(MIG, 'w', encoding='utf-8', newline='\n').write(sql)

# ---------------------------------------------------------------- outputs: review workbook
OWNER = [
    ('HOOKKA sofas', f'HOOKA.pdf cuts every sofa list short: e.g. 5530 says "13 SKUs" but prints only 1A(LHF), 1A(RHF), 1NA, 1S, 2A(LHF). '
                     f'{n_sofa} sofa SKUs were created from what it prints; the catalogue counts {sum(m["count"] for m in models if m["category"] == "SOFA")}. '
                     'Please get HOOKKA\'s full compartment list per model (or say if we may take it from Houzs\'s own records, since HOOKKA is Houzs\'s factory). '
                     'A missing compartment can also be ticked on the model page (Products > model > Allowed options); the engine then creates its SKU, but without the HOOKKA link.',
                     'HOOKA.pdf p4-p5'),
    ('HOOKKA bedframes', f'Same for 6 bedframe models (1003, 1007, 2006, 2008, 2009, DIVAN): the size list is cut off. {n_bed} bedframe SKUs were created from the printed sizes; '
                         'e.g. 1003 has no 6FT SKU because the page stops at "5FT…". Also 1019, 1023, 1030, 2023, 2033, 2038 have more SKUs than sizes (1019: "8 SKUs · 3.5FT, 3FT, 5FT, 6FT") - what is the other choice (height? storage?).',
                         'HOOKA.pdf p1-p3'),
    ('HOOKKA cm sizes', 'Sizes HOOKKA prints in cm other than 200x200 (e.g. 152x200, 220x220) were created as their own SKU under the engine\'s CUSTOM size (SP). OK, or should any be a standard size?', 'HOOKA.pdf p1-p3'),
    ('5543 "12" X 28""', 'Sofa 5543\'s list starts with 12" X 28" - the long pillow size, not a sofa compartment. Not created; the pillow is SKU ACC-HKA-LONGPILLOW. Correct?', 'HOOKA.pdf p4'),
    ('HOOKKA unit', 'The catalogue writes no unit. Bedframes / sofas stay "Unit", pillows "Piece". OK?', 'HOOKA.pdf'),
    ('LB stone unit', f'{sum(1 for a in audit if a["q"] == "lb")} LB quartz / sintered stone SKUs: the catalogues give the slab size only, no unit. Now "Piece" (= one slab). Bought per slab, per foot run or per square foot?', 'LB catalogues'),
    ('DHI flooring unit', f'{sum(1 for a in audit if a["q"] == "dhi_floor")} DHI SPC flooring SKUs: no unit or price, only packing "24.2 sqft (10pcs/ctn)". Now "Square foot". Per square foot or per carton?', 'DHI_Catalogue V3.pdf'),
    ('DHI accessories unit', f'{sum(1 for a in audit if a["q"] == "dhi_acc")} DHI skirting / moulding SKUs: photo pages with no unit. Now "Piece" (one length). OK?', 'DHI Accessories Showcase'),
    ('K & K lines without a unit', 'K & K prints no unit on these lines, kept as they were: ' + '; '.join(
        f'{a["code"]} {a["name"]} = {a["old"]}' for a in audit if a['q'] == 'kk'), 'K & K PRICE LIST.pdf p3-p4'),
    ('UNILOC stair nosing', 'ORI SPC stair nosing 5mm / 6mm: RM39 / RM42 is printed in the "RM/FT²" column (4pcs, no carton price), so the unit is now "Square foot" as written. If UNILOC means per piece, say so.', '0022_UNILOC Price List p1'),
]
OWNER = [o for o in OWNER if o[0] != 'UNILOC stair nosing' or any(a['q'] == 'nosing' for a in audit)]

out = openpyxl.Workbook()
bold = Font(bold=True); head_fill = PatternFill('solid', fgColor='DDE6F0')
def sheet(title, header, data, widths):
    ws = out.create_sheet(title)
    ws.append(header)
    for c in ws[1]: c.font = bold; c.fill = head_fill
    for r in data: ws.append(list(r))
    for i, w in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    return ws
out.remove(out.active)
conf_order = {'high': 0, 'medium': 1, 'none': 2}
uname = dict((c, l) for c, l, _ in UOM_LABELS); uname.update({c: l for c, l, _, _ in NEW_UOMS})
sheet('Unit changes', ['SKU', 'Name', 'Supplier', 'Old unit', 'New unit', 'Changed?', 'Evidence (file + page + wording)', 'Confidence'],
      [(a['code'], a['name'], a['supplier'], a['old'], a['new'], 'Y' if a['new'] != a['old'] else '', a['evidence'],
        'no unit on document - kept, see Owner to decide' if a['confidence'] == 'none' else a['confidence'])
       for a in sorted(audit, key=lambda a: (a['new'] == a['old'], conf_order[a['confidence']], a['code']))],
      [30, 46, 26, 9, 9, 9, 80, 22])
sheet('Owner to decide', ['#', 'Topic', 'Question', 'Where'], [(i + 1, *o) for i, o in enumerate(OWNER)], [4, 22, 110, 26])
sheet('Bedframe sizes', ['New SKU', 'Name', 'Model (base_model)', 'Size code', 'Size label', 'Printed as', 'HOOKKA code', 'Page', 'Catalogue line', 'Replaces'],
      [(n['code'], n['name'], n['base_model'], n['size_code'], n['size_label'], n['printed'], n['supplier_sku'], f'HOOKA.pdf p{n["page"]}',
        next(m['printed'] for m in models if m['model_code'] == n['base_model']), n['old_code']) for n in new_rows if n['category'] == 'BEDFRAME'],
      [30, 36, 16, 9, 12, 14, 18, 12, 52, 16])
sheet('Sofa compartments', ['New SKU', 'Name', 'Model (base_model)', 'Compartment', 'HOOKKA code', 'Page', 'Catalogue line', 'Replaces'],
      [(n['code'], n['name'], n['base_model'], n['printed'], n['supplier_sku'], f'HOOKA.pdf p{n["page"]}',
        next(m['printed'] for m in models if m['model_code'] == n['base_model']), n['old_code']) for n in new_rows if n['category'] == 'SOFA'],
      [26, 36, 16, 12, 16, 12, 52, 16])
sheet('Models', ['Model', 'Category', 'Catalogue line', 'Page', 'SKUs in catalogue', 'SKUs created', 'List cut off?', 'Printed but not created'],
      [(m['model_code'], m['category'], m['printed'], f'HOOKA.pdf p{m["page"]}', m['count'], m['created'], 'Y' if m['cut'] else '', ', '.join(m['skipped']))
       for m in models], [16, 10, 56, 12, 10, 10, 10, 20])
sheet('Deleted', ['SKU', 'Name', 'Supplier item code', 'Source', 'What happens'],
      [(c, by_code[c]['name'], by_code[c]['supplier_item_code'], short(by_code[c]['item_source']),
        'Deleted with its HOOKKA supplier link, cost history (none) and model, after checking no document, stock or price row names it (else deactivated); written to the audit log')
       for c in DELETE_2990]
      + [(c, by_code[c]['name'], by_code[c]['supplier_item_code'], short(by_code[c]['item_source']),
          'Replaced by the size / compartment SKUs above; deleted the same way (else deactivated)') for c in REPLACED],
      [24, 34, 16, 30, 90])
sheet('Units', ['Code', 'Old label', 'Old print label', 'New label', 'New print label', 'Why'],
      [(c, *OLD_LABELS[c], l, p, 'clearer label') for c, l, p in UOM_LABELS if OLD_LABELS[c] != (l, p)]
      + [(c, '', '', l, p, 'added: K & K writes "no" / "nos" (doors, silicone, door closer)') for c, l, p, _ in NEW_UOMS],
      [8, 16, 14, 24, 14, 60])
rm = out.create_sheet('Read me', 0)
for line in [
    'Catalog fix, 21 Oct 2026 - what migration 0132 changes in the engine catalog (first import of 20 Oct 2026).',
    '',
    f'Unit changes: {len(audit)} SKUs checked against their supplier document; {len(changes)} change unit (Changed? = Y). Rule: the unit the supplier writes.',
    f'  Where a document writes no unit the SKU keeps its unit and is on "Owner to decide" ({sum(1 for a in audit if a["confidence"] == "none")} SKUs).',
    f'Bedframe sizes: {n_bed} SKUs for {sum(1 for m in models if m["category"] == "BEDFRAME")} HOOKKA bedframe models. Sofa compartments: {n_sofa} SKUs for {sum(1 for m in models if m["category"] == "SOFA")} sofa models.',
    '  Only sizes / compartments HOOKA.pdf prints were created; the catalogue cuts several lists short (sheet Models).',
    f'Deleted: the 13 "2990 list" sofas and the {len(REPLACED)} single bedframe / sofa SKUs the new ones replace.',
    'Units: clearer labels (FT is now "Foot run"), unit NOS added. Square inch / metre / square metre were not added: no supplier document uses them.',
    'Costs stay empty for HOOKKA: the catalogue has no prices.',
]:
    rm.append([line])
rm.column_dimensions['A'].width = 140
out.save(os.path.join(DOCS, 'catalog-fix-review.xlsx'))

print(f'units checked {len(audit)}, changed {len(changes)}: ' + ', '.join(f'{o}->{n} {c}' for (o, n), c in
      sorted(__import__("collections").Counter((a["old"], a["new"]) for a in changes).items())))
print(f'owner-to-decide SKUs {sum(1 for a in audit if a["confidence"] == "none")}; bedframe SKUs {n_bed}; sofa SKUs {n_sofa}; models {len(models)}; '
      f'skipped {[(m["model_code"], m["skipped"]) for m in models if m["skipped"]]}')
