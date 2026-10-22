"""HOOKKA quotation 300-T (22 Oct 2026) as the HOOKKA catalog, plus the AKEMI mattresses from Houzs Century's list.

    python -I docs/catalog-import/build/hookka.py <repo root>

Reads build/hookka_300t.json (hookka_parse.py), build/akemi_mattress.csv (the 231 AKEMI rows of Houzs Century's
mattress list of 22 Oct 2026), catalog-fix-data.json (the HOOKKA SKUs migration 0132 created) and catalog-data.json
(the 3 HOOKKA pillows of the first import), and migration 0129's master maintenance config.
Writes docs/catalog-import/hookka-300t-data.json, docs/catalog-import/hookka-300t-review.xlsx and
migrations-pg/0142_hookka_quote_akemi_mattresses.sql. Edit this script, not the migration.

Engine shape (engine/backend/src/scm/lib/cost-anchor-sync.ts, shared/mfg-pricing.ts): on a SKU, base_price_sen is the
PRICE_2 cost, price1_sen the PRICE_1 cost, seat_height_prices[].priceSen the sofa cost per (seat height, tier); the
supplier link carries the same as price_matrix ({P2,P1} for a bedframe, {"<seat>": {P2,P3}} for a sofa) plus a flat
unit_price_sen. The link is the cost anchor, so an edit on either side keeps the other in step.
"""
import sys, os, re, json, csv, hashlib
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

ROOT = sys.argv[1]
DOCS = os.path.join(ROOT, 'docs', 'catalog-import')
BUILD = os.path.join(DOCS, 'build')
MIG = os.path.join(ROOT, 'migrations-pg', '0142_hookka_quote_akemi_mattresses.sql')
TAG = 'hookka-300t 2026-10-22'
EFF = '2026-10-22'
HK_NOTE = 'HOOKKA quotation 300-T 22 Oct 2026'
AK_NOTE = 'Houzs Century mattress list 2026-10-22'
HK_BRANDING = 'HOOKKA INDUSTRIES SDN BHD'

Q = json.load(open(os.path.join(BUILD, 'hookka_300t.json'), encoding='utf-8'))
FIX = json.load(open(os.path.join(DOCS, 'catalog-fix-data.json'), encoding='utf-8'))
FIRST = json.load(open(os.path.join(DOCS, 'catalog-data.json'), encoding='utf-8'))
OLD = {n['code']: n for n in FIX['new_skus']}
OLD.update({k['code']: k for k in FIRST['skus'] if k['code'].startswith('ACC-HKA-')})
assert len(OLD) == 163, len(OLD)

def sku_id(code):
    return 'mfg-' + hashlib.sha256(code.encode('utf-8')).hexdigest()[:12]

def rm(sen):
    return None if sen is None else round(sen / 100, 2)

skus, models = [], {}

# ---------------------------------------------------------------- bedframes (161)
FT = {'6FT': 'K', '5FT': 'Q', '3FT': 'S', '3.5FT': 'SS'}
for b in Q['bedframes']:
    # "1030-(HF)(W)(SS)" is HOOKKA's spelling of 1030(HF)(W)-(SS) (its siblings are 1030(HF)(W)-(K) ...).
    norm = re.sub(r'^(\d+)-(\(HF\)\(W\))(\(\w+\))$', lambda g: g.group(1) + g.group(2) + '-' + g.group(3), b['code'])
    m = re.fullmatch(r'(.+?)-\(?([A-Z0-9X]+)\)?(?:\s*\((\d+X\d+)CM\))?', norm)
    assert m, b['code']
    model, token, extra = m.groups()
    if token == 'SP' and extra and '(' not in b['code'].split('-', 1)[1][:1]:
        token = extra                                   # "2008(A)-SP (210X240CM)": the cm size is the real code
    size = b['size']
    if token in ('K', 'Q', 'S', 'SS', 'SK', 'SP'):
        size_code = token
    elif size == '200CMX200CM' or size.upper() == 'SUPERKING':
        size_code = 'SK'
    else:
        size_code = 'SP'
    size_label = '200X200CM' if size.upper() == 'SUPERKING' else size.replace('CMX', 'X')
    p2, p1 = b['price2_sen'], b['price1_sen']
    base = f'BED-HKA-{model}'
    code = f'{base}-({token})'
    mname = re.sub(r'\s*\((\d+(\.\d+)?FT|\d+X\d+CM|SUPERKING)\)', '', b['description']).strip()
    models.setdefault(base, {'model_code': base, 'name': mname, 'category': 'BEDFRAME', 'branding': HK_BRANDING,
                             'allowed_options': {'sizes': []}})
    if size_code not in models[base]['allowed_options']['sizes']:
        models[base]['allowed_options']['sizes'].append(size_code)
    flat = p2 if p2 is not None else p1
    skus.append({
        'family': 'BEDFRAME', 'id': sku_id(code), 'code': code, 'name': b['description'], 'category': 'BEDFRAME',
        'sub_category': 'BEDFRAME', 'uom': 'UNIT', 'branding': HK_BRANDING, 'base_model': base, 'size_code': size_code,
        'size_label': size_label, 'supplier': 'HOOKKA', 'supplier_sku': b['code'],
        'description': f'HOOKKA {b["code"]}: {b["description"]}. Quotation 300-T, 22 Oct 2026, p{b["page"]}.',
        'base_price_sen': p2, 'price1_sen': p1, 'seat_height_prices': None, 'cost_price_sen': flat or 0,
        'unit_price_sen': flat or 0, 'price_matrix': {k: v for k, v in (('P2', p2), ('P1', p1)) if v is not None} or None,
        'priced': flat is not None, 'page': b['page'], 'quote_p2': p2, 'quote_p1': p1,
    })

# ---------------------------------------------------------------- sofas (one SKU per compartment; LHF/RHF split)
TIER = {'P1': 'PRICE_1', 'P2': 'PRICE_2', 'P3': 'PRICE_3'}
sofa = {}
for r in Q['sofa_rows']:
    comp = r['compartment'].replace(' ', '')
    sides = [comp.replace('(LHF/RHF)', '(LHF)'), comp.replace('(LHF/RHF)', '(RHF)')] if '(LHF/RHF)' in comp else [comp]
    for c in sides:
        k = (r['model'], c)
        sofa.setdefault(k, {'model': r['model'], 'comp': c, 'printed': r['compartment'], 'page': r['page'], 'tiers': {}})
        sofa[k]['tiers'][r['tier']] = r['prices']
for (mdl, comp), s in sofa.items():
    base = f'SOF-HKA-{mdl}'
    code = f'{base}-{comp}'
    models.setdefault(base, {'model_code': base, 'name': f'HOOKKA Sofa {mdl}', 'category': 'SOFA', 'branding': HK_BRANDING,
                             'allowed_options': {'compartments': []}})
    models[base]['allowed_options']['compartments'].append(comp)
    heights = list(next(iter(s['tiers'].values())).keys())
    matrix, seat = {}, []
    for h in heights:
        cell = {t: s['tiers'][t][h] for t in ('P2', 'P3', 'P1') if t in s['tiers'] and s['tiers'][t].get(h) is not None}
        if cell:
            matrix[h] = cell
            for t, v in cell.items():
                seat.append({'height': h, 'tier': TIER[t], 'priceSen': v})
    # Flat cost (a line with no seat height, e.g. a renovation contract line): P2 at the smallest priced seat.
    flat = next((matrix[h]['P2'] for h in heights if h in matrix and 'P2' in matrix[h]), None)
    if flat is None:
        flat = next((v for h in heights if h in matrix for v in matrix[h].values()), None)
    skus.append({
        'family': 'SOFA', 'id': sku_id(code), 'code': code, 'name': f'SOFA {mdl} {comp}', 'category': 'SOFA',
        'sub_category': 'SOFA', 'uom': 'UNIT', 'branding': HK_BRANDING, 'base_model': base, 'size_code': None,
        'size_label': None, 'supplier': 'HOOKKA', 'supplier_sku': f'{mdl} {comp}',
        'description': f'HOOKKA sofa {mdl}, compartment {comp} (quoted as "SOFA {mdl} {s["printed"]}"). Cost per seat height '
                       f'{", ".join(h + chr(34) for h in heights)} and fabric tier {"/".join(sorted(s["tiers"]))}. '
                       f'Quotation 300-T, 22 Oct 2026, p{s["page"]}.',
        'base_price_sen': flat, 'price1_sen': None, 'seat_height_prices': seat, 'cost_price_sen': flat or 0,
        'unit_price_sen': flat or 0, 'price_matrix': matrix or None, 'priced': flat is not None, 'page': s['page'],
        'heights': heights, 'tiers': s['tiers'],
    })

# ---------------------------------------------------------------- accessories (5); the 3 pillows keep their codes
ACC_CODE = {'LONG PILLOW': 'ACC-HKA-LONGPILLOW', 'SQUARE PILLOW': 'ACC-HKA-SQUAREPILLOW',
            'SQUARE PILLOW SP': 'ACC-HKA-SQUAREPILLOWSP', '5543-HDRST': 'ACC-HKA-5543-HDRST',
            '5543-LONG PILLOW': 'ACC-HKA-5543-LONGPILLOW'}
for a in Q['accessories']:
    code = ACC_CODE[a['code']]
    old = OLD.get(code)
    skus.append({
        'family': 'ACCESSORY', 'id': sku_id(code), 'code': code,
        'name': old['name'] if old else 'HOOKKA ' + a['description'].title().replace('"X', '"x'),
        'category': 'FURNITURE', 'sub_category': 'ACCESSORY', 'uom': 'PCS', 'branding': HK_BRANDING, 'base_model': None,
        'size_code': None, 'size_label': None, 'supplier': 'HOOKKA', 'supplier_sku': a['code'],
        'description': f'HOOKKA {a["code"]}: {a["description"]}, size {a["size"]}. Quotation 300-T, 22 Oct 2026, p{a["page"]}.',
        'base_price_sen': None, 'price1_sen': None, 'seat_height_prices': None, 'cost_price_sen': a['price_sen'],
        'unit_price_sen': a['price_sen'], 'price_matrix': None, 'priced': True, 'page': a['page'],
    })

# ---------------------------------------------------------------- AKEMI mattresses (231) from Houzs Century's list
SZ = {'6FT': 'K', '5FT': 'Q', '3.5FT': 'SS', '3FT': 'S', '200X200CM': 'SK', 'CUSTOM': 'SP'}
akemi = list(csv.DictReader(open(os.path.join(BUILD, 'akemi_mattress.csv'), encoding='utf-8')))
assert len(akemi) == 231 and all(r['code'].startswith('AKEMI') for r in akemi)
for r in akemi:
    m = re.fullmatch(r'AKEMI (.+) MATT \((K|Q|S|SS|SK|SP)\)', r['code'])
    assert m, r['code']
    mcode, sz = m.groups()
    assert SZ[r['size_label']] == sz, r
    base_p = int(round(float(r['base_price']) * 100)) if r['base_price'] else None
    p1 = int(round(float(r['price1']) * 100)) if r['price1'] else None
    mdl = models.setdefault(('AKEMI', mcode), {'model_code': mcode, 'name': mcode, 'category': 'MATTRESS', 'branding': 'AKEMI',
                                               'allowed_options': {'sizes': []}})
    mdl['allowed_options']['sizes'].append(sz)
    th = re.search(r'X(\d+)CM\)$', r['name'])
    if th:
        mdl.setdefault('_thick', set()).add(int(th.group(1)))
    skus.append({
        'family': 'MATTRESS', 'id': sku_id(r['code']), 'code': r['code'], 'name': r['name'], 'category': 'MATTRESS',
        'sub_category': 'MATTRESS', 'uom': 'UNIT', 'branding': 'AKEMI', 'base_model': mcode, 'size_code': sz,
        'size_label': r['size_label'], 'supplier': 'HOUZS-CENTURY', 'supplier_sku': r['code'],
        'description': f'AKEMI {mcode} mattress, size {r["size_label"]}. Houzs Century mattress list, 22 Oct 2026.',
        'base_price_sen': base_p, 'price1_sen': p1, 'seat_height_prices': None, 'cost_price_sen': base_p or 0,
        'unit_price_sen': base_p or 0, 'price_matrix': None, 'priced': base_p is not None,
    })
for mdl in models.values():
    t = mdl.pop('_thick', None)
    if t and len(t) == 1:
        mdl['allowed_options']['mattress_thickness_cm'] = t.pop()

codes = [s['code'] for s in skus]
assert len(codes) == len(set(codes)), [c for c in codes if codes.count(c) > 1]
fam = {f: sum(1 for s in skus if s['family'] == f) for f in ('BEDFRAME', 'SOFA', 'ACCESSORY', 'MATTRESS')}
assert fam['BEDFRAME'] == 161 and fam['ACCESSORY'] == 5 and fam['MATTRESS'] == 231, fam
for s in skus:
    s['note'] = (f'{TAG}: {HK_NOTE if s["supplier"] == "HOOKKA" else AK_NOTE}'
                 + ('' if s['priced'] else '; PRICE NEEDED - no price on the list'))
    s['exists'] = s['code'] in OLD
retired = sorted(c for c in OLD if c not in set(codes))

# ---------------------------------------------------------------- surcharges: HOOKKA's price book vs the master (0129)
m129 = open(os.path.join(ROOT, 'migrations-pg', '0129_engine_houzs_product_maintenance.sql'), encoding='utf-8').read()
mj = re.search(r"\|\| '(\{.*?\})'::jsonb", m129, re.S).group(1).replace("''", "'")
MASTER = json.loads(mj)
POOLS = {'Bedframe/Divan Heights': 'divanHeights', 'Bedframe/Total Heights': 'totalHeights', 'Bedframe/Leg Heights': 'legHeights',
         'Bedframe/Specials': 'specials', 'Sofa/Leg Heights': 'sofaLegHeights', 'Sofa/Specials': 'sofaSpecials'}
supplier_pools, surch_rows = {}, []
for src, key in POOLS.items():
    pool = []
    master = {e['value']: e for e in MASTER.get(key, [])}
    for it in Q['surcharges'][src]:
        val = 'No Leg' if it['item'] == 'No Leg' else it['item']
        sen = it['surcharge_sen'] or 0
        entry = {'value': val, 'priceSen': sen}
        if val in master and 'packSeparately' in master[val]:
            entry['packSeparately'] = master[val]['packSeparately']
        pool.append(entry)
        mv = master.get(val)
        surch_rows.append([src.split('/')[0], src.split('/')[1], val, rm(sen) if it['surcharge_sen'] is not None else '— (0)',
                           rm(mv['priceSen']) if mv else 'not in master',
                           'same' if mv and mv['priceSen'] == sen else ('only HOOKKA' if not mv else 'DIFFERENT')])
    for v, mv in master.items():
        if v not in [e['value'] for e in pool]:
            surch_rows.append([src.split('/')[0], src.split('/')[1], v, 'not quoted', rm(mv.get('priceSen', 0)),
                               'only master (kept in HOOKKA price book at master value)'])
    supplier_pools[key] = pool
gaps = Q['surcharges']['Bedframe/Gaps']
supplier_pools['gaps'] = gaps
sizes = [s.strip() for s in Q['surcharges']['Sofa/Sizes']]
for g in gaps:
    surch_rows.append(['Bedframe', 'Gaps', g, 'listed', 'listed' if g in MASTER['gaps'] else 'not in master',
                       'same' if g in MASTER['gaps'] else 'only HOOKKA'])
for z in sizes:
    surch_rows.append(['Sofa', 'Sizes', z, 'listed', 'listed' if z in MASTER['sofaSizes'] else 'not in master',
                       'same' if z in MASTER['sofaSizes'] else 'only HOOKKA (not added: not a seat height any SKU is priced at)'])

json.dump({'tag': TAG, 'effective': EFF, 'skus': skus, 'models': list(models.values()), 'retired_candidates': retired,
           'supplier_pools': supplier_pools}, open(os.path.join(DOCS, 'hookka-300t-data.json'), 'w', encoding='utf-8'),
          indent=1, ensure_ascii=False)

# ---------------------------------------------------------------- customer-supplied placeholders (owner, 22 Oct 2026)
# One "Customer <sub-category>" SKU per active sub-category (read from the database when the migration runs, so a
# sub-category added on production is covered too) plus one per group, all on supplier CUSTOMER at cost 0. The group
# SKU sits under the group's most general sub-category (every group but Service requires one).
SUBCATS = json.load(open(os.path.join(BUILD, 'subcategories-2026-10-22.json'), encoding='utf-8'))
CUS_GROUPS = [  # group, SKU code, name, sub-category it sits under
    ('WORKS', 'CUS-WORKS', 'Customer Works', 'PROTECTION'), ('BUILTIN', 'CUS-BUILTIN', 'Customer Built-in', 'SHELVING'),
    ('FURNITURE', 'CUS-FURNITURE', 'Customer Furniture', 'LOOSE_STORAGE'), ('LIGHTING', 'CUS-LIGHTING', 'Customer Lighting', 'DOWNLIGHT'),
    ('APPLIANCE', 'CUS-APPLIANCE', 'Customer Appliance', 'TV'), ('SOFT', 'CUS-SOFT', 'Customer Soft furnishing', 'DECORATION'),
    ('SERVICE', 'CUS-SERVICE', 'Customer Service', None)]
CUS_NOTE = 'customer-supplied 2026-10-22: the customer orders this item, The Conts manages it; cost 0 by design (not price needed)'
CUS_DESC = 'Customer-supplied: the customer orders / buys this item and The Conts manages it. Cost 0; no purchase order is raised.'

# ---------------------------------------------------------------- migration
def lit(obj):
    s = json.dumps(obj, ensure_ascii=False)
    assert '$hq$' not in s
    return f'$hq${s}$hq$'

SKU_COLS = ['id', 'code', 'name', 'category', 'sub_category', 'uom', 'branding', 'base_model', 'size_code', 'size_label',
            'supplier', 'supplier_sku', 'description', 'base_price_sen', 'price1_sen', 'seat_height_prices', 'cost_price_sen',
            'unit_price_sen', 'price_matrix', 'priced', 'note']
sku_rows = [{k: s[k] for k in SKU_COLS} for s in skus]
CO = "(select id from public.companies where code = 'HOUZS')"
sql = f"""-- REVERSAL: none in SQL. Restore the HOOKKA SKUs from migration 0132 (docs/catalog-import/catalog-fix-data.json) and
--   delete the SKUs, models and supplier links named in docs/catalog-import/hookka-300t-data.json that 0132 did not create;
--   delete scm.maintenance_config_history 'mch-hookka-300t-20261022' and the price-history rows noted '{TAG}%'.
-- HOOKKA's quotation to The Conts (customer 300-T, effective 22 Oct 2026) becomes the HOOKKA catalog, and the AKEMI
-- mattresses from Houzs Century's list of 22 Oct 2026 are added (owner, 22 Oct 2026). Review: docs/catalog-import/hookka-300t-review.xlsx.
--  1. Suppliers HOOKKA and HOUZS CENTURY SDN BHD (found by code / name first, created only when missing).
--  2. Models: one BEDFRAME model per HOOKKA family (1003, 1003(A), 1003(A)(HF)(W), ...) with its sizes, one SOFA model per
--     sofa with its compartments, one MATTRESS model per AKEMI range (branding AKEMI, the engine's "AKEMI <model> MATT (size)" code).
--  3. SKUs: {fam['BEDFRAME']} bedframes (BED-HKA-<HOOKKA model>-(<size>)), {fam['SOFA']} sofa compartments (SOF-HKA-<model>-<compartment>,
--     "(LHF/RHF)" split into the two sides the engine's compartments use), {fam['ACCESSORY']} accessories, {fam['MATTRESS']} AKEMI mattresses.
--     A SKU that exists is updated in place (code, id and status kept), the rest are created Active.
--     Cost the engine's way (lib/cost-anchor-sync.ts): base_price_sen = Price 2 cost, price1_sen = Price 1 cost,
--     seat_height_prices = sofa cost per seat height and tier (P2 -> PRICE_2, P3 -> PRICE_3), cost_price_sen = flat cost.
--  4. Each SKU has exactly one main supplier link (HOOKKA / Houzs Century): supplier item code = the list's code, flat
--     unit_price_sen, price_matrix ({{P2,P1}} bedframe, {{"<seat>": {{P2,P3}}}} sofa), cost anchor, valid from {EFF};
--     one supplier_binding_price_history row per priced link. Unpriced rows (39 AKEMI) are linked with cost 0, "PRICE NEEDED".
--  5. HOOKKA SKUs from 0132 / the first import that the quotation does not contain are deleted with their supplier link,
--     cost history and model when nothing else names them, else deactivated; each written to scm.entity_audit_log.
--  7. Customer-supplied placeholder SKUs ("Customer <sub-category>", one per active sub-category and one per group) on
--     supplier CUSTOMER at cost 0; a renovation contract raises no PO for them.
--  6. HOOKKA's own surcharges (divan / total / leg heights, specials, gaps) become HOOKKA's supplier price book
--     (maintenance_config_history scope 'supplier:<HOOKKA id>'), read by the PO cost; the master config is not changed.
-- Built by docs/catalog-import/build/hookka.py; edit that script, not this file. Safe to run again: it finds nothing to change.
set local search_path = scm, public, extensions;

-- 1. Suppliers.
insert into scm.suppliers (company_id, code, name, status, notes)
select {CO}, 'HOOKKA', 'HOOKKA INDUSTRIES SDN BHD', 'ACTIVE', 'Furniture supplier (catalog import)'
 where exists (select 1 from public.companies where code = 'HOUZS')
   and not exists (select 1 from scm.suppliers s where s.company_id = {CO} and (s.code = 'HOOKKA' or s.name ilike 'HOOKKA INDUSTRIES%'));
insert into scm.suppliers (company_id, code, name, status, notes)
select {CO}, 'HOUZS-CENTURY', 'HOUZS CENTURY SDN BHD', 'ACTIVE', 'Mattress supplier (AKEMI); {AK_NOTE}'
 where exists (select 1 from public.companies where code = 'HOUZS')
   and not exists (select 1 from scm.suppliers s where s.company_id = {CO} and (s.code = 'HOUZS-CENTURY' or s.name ilike 'HOUZS CENTURY%'));

create temp table hq_sup on commit drop as
select 'HOOKKA'::text as key, (select s.id from scm.suppliers s where s.company_id = {CO} and (s.code = 'HOOKKA' or s.name ilike 'HOOKKA INDUSTRIES%')
                                order by (s.code = 'HOOKKA') desc, s.created_at limit 1) as supplier_id
union all
select 'HOUZS-CENTURY', (select s.id from scm.suppliers s where s.company_id = {CO} and (s.code = 'HOUZS-CENTURY' or s.name ilike 'HOUZS CENTURY%')
                          order by (s.code = 'HOUZS-CENTURY') desc, s.created_at limit 1);

create temp table hq_sku on commit drop as
select n.*, s.supplier_id from jsonb_to_recordset({lit(sku_rows)}::jsonb) as n(
  id text, code text, name text, category text, sub_category text, uom text, branding text, base_model text, size_code text,
  size_label text, supplier text, supplier_sku text, description text, base_price_sen int, price1_sen int,
  seat_height_prices jsonb, cost_price_sen int, unit_price_sen int, price_matrix jsonb, priced boolean, note text)
  join hq_sup s on s.key = n.supplier
 where s.supplier_id is not null;

create temp table hq_models on commit drop as
select * from jsonb_to_recordset({lit(list(models.values()))}::jsonb) as m(model_code text, name text, category text, branding text, allowed_options jsonb);

-- 2. Models (one per family; an existing one keeps its id and gets the quotation's sizes / compartments).
insert into scm.product_models (company_id, model_code, name, category, branding, allowed_options)
select {CO}, x.model_code, x.name, x.category::scm.mfg_product_category, x.branding, x.allowed_options from hq_models x
 where exists (select 1 from hq_sku)
on conflict (company_id, model_code, category) do update
   set name = excluded.name, branding = excluded.branding,
       allowed_options = scm.product_models.allowed_options || excluded.allowed_options, updated_at = now()
 where (scm.product_models.name, scm.product_models.branding, scm.product_models.allowed_options)
       is distinct from (excluded.name, excluded.branding, scm.product_models.allowed_options || excluded.allowed_options);

-- 3. SKUs: update the ones that exist, create the rest.
update scm.mfg_products p
   set name = n.name, category = n.category::scm.mfg_product_category, sub_category = n.sub_category, branding = n.branding,
       base_model = n.base_model, size_code = n.size_code, size_label = n.size_label, description = n.description,
       base_price_sen = n.base_price_sen, price1_sen = n.price1_sen, seat_height_prices = n.seat_height_prices,
       cost_price_sen = n.cost_price_sen, model_id = m.id, updated_at = now()
  from hq_sku n
  join scm.product_models m on m.company_id = {CO} and m.model_code = n.base_model and m.category::text = n.category
 where p.company_id = {CO} and p.code = n.code
   and (p.name, p.category::text, p.sub_category, p.branding, p.base_model, p.size_code, p.size_label, p.description,
        p.base_price_sen, p.price1_sen, p.seat_height_prices, p.cost_price_sen, p.model_id)
       is distinct from (n.name, n.category, n.sub_category, n.branding, n.base_model, n.size_code, n.size_label, n.description,
        n.base_price_sen, n.price1_sen, n.seat_height_prices, n.cost_price_sen, m.id);
update scm.mfg_products p
   set name = n.name, sub_category = n.sub_category, description = n.description, cost_price_sen = n.cost_price_sen, updated_at = now()
  from hq_sku n
 where n.base_model is null and p.company_id = {CO} and p.code = n.code
   and (p.name, p.sub_category, p.description, p.cost_price_sen) is distinct from (n.name, n.sub_category, n.description, n.cost_price_sen);

insert into scm.mfg_products (id, company_id, code, name, category, sub_category, uom, status, cost_price_sen, base_price_sen,
                              price1_sen, seat_height_prices, description, branding, model_id, base_model, size_code, size_label)
select n.id, {CO}, n.code, n.name, n.category::scm.mfg_product_category, n.sub_category, n.uom, 'ACTIVE', n.cost_price_sen,
       n.base_price_sen, n.price1_sen, n.seat_height_prices, n.description, n.branding, m.id, n.base_model, n.size_code, n.size_label
  from hq_sku n
  left join scm.product_models m on m.company_id = {CO} and m.model_code = n.base_model and m.category::text = n.category
 where not exists (select 1 from scm.mfg_products p where p.company_id = {CO} and p.code = n.code)
on conflict do nothing;
-- An accessory with no family model gets its own model, as the API and the first import do.
insert into scm.product_models (company_id, model_code, name, category, branding)
select {CO}, n.code, n.name, 'FURNITURE', n.branding from hq_sku n
 where n.base_model is null
on conflict (company_id, model_code, category) do nothing;
update scm.mfg_products p set model_id = m.id
  from hq_sku n join scm.product_models m on m.company_id = {CO} and m.model_code = n.code and m.category::text = 'FURNITURE'
 where n.base_model is null and p.company_id = {CO} and p.code = n.code and p.model_id is null;

-- 4. One main supplier link per SKU, carrying the quoted cost; any other link for the SKU stops being main.
update scm.supplier_material_bindings b
   set is_main_supplier = false, is_cost_anchor = false, updated_at = now()
  from hq_sku n
 where b.company_id = {CO} and b.material_kind = 'mfg_product' and b.item_code = n.code and b.supplier_id <> n.supplier_id
   and (b.is_main_supplier or b.is_cost_anchor);
update scm.supplier_material_bindings b
   set material_name = n.name, supplier_sku = n.supplier_sku, unit_price_sen = n.unit_price_sen, price_matrix = n.price_matrix,
       price_valid_from = case when n.priced then date '{EFF}' end, is_main_supplier = true, is_cost_anchor = n.priced,
       notes = n.note, updated_at = now()
  from hq_sku n
 where b.company_id = {CO} and b.material_kind = 'mfg_product' and b.item_code = n.code and b.supplier_id = n.supplier_id
   and (b.material_name, b.supplier_sku, b.unit_price_sen, b.price_matrix, b.price_valid_from, b.is_main_supplier, b.is_cost_anchor, b.notes)
       is distinct from (n.name, n.supplier_sku, n.unit_price_sen, n.price_matrix, case when n.priced then date '{EFF}' end, true, n.priced, n.note);
insert into scm.supplier_material_bindings
    (company_id, supplier_id, material_kind, item_code, material_name, supplier_sku, unit_price_sen, price_matrix,
     price_valid_from, is_main_supplier, is_cost_anchor, notes)
select {CO}, n.supplier_id, 'mfg_product', n.code, n.name, n.supplier_sku, n.unit_price_sen, n.price_matrix,
       case when n.priced then date '{EFF}' end, true, n.priced, n.note
  from hq_sku n
 where not exists (select 1 from scm.supplier_material_bindings b where b.company_id = {CO} and b.material_kind = 'mfg_product'
                                                                    and b.item_code = n.code and b.supplier_id = n.supplier_id);
insert into scm.supplier_binding_price_history
    (company_id, supplier_id, material_kind, item_code, unit_price_sen, price_matrix, is_main_supplier, effective_from, notes,
     created_by, applied_at)
select {CO}, n.supplier_id, 'mfg_product', n.code, n.unit_price_sen, n.price_matrix, true, date '{EFF}', n.note, 'migration 0142', now()
  from hq_sku n
 where n.priced
   and not exists (select 1 from scm.supplier_binding_price_history h where h.company_id = {CO} and h.material_kind = 'mfg_product'
                     and h.item_code = n.code and h.supplier_id = n.supplier_id and h.effective_from = date '{EFF}');

-- 5. HOOKKA SKUs the quotation does not contain.
create temp table hq_retire on commit drop as
select p.id, p.code, p.model_id, 0::bigint as refs, ''::text as used_in
  from scm.mfg_products p
 where p.company_id = {CO} and p.code = any({lit(retired).replace('$hq$[', "$hq${").replace(']$hq$', '}$hq$')}::text[])
   and exists (select 1 from hq_sku)
   and not exists (select 1 from scm.entity_audit_log a where a.entity_type = 'mfg_product' and a.entity_id = p.code
                                                        and a.note like '{TAG}%');
do $$
declare t record;
begin
  for t in
    select c.table_schema as s, c.table_name as tb, c.column_name as col
      from information_schema.columns c
      join information_schema.tables x on x.table_schema = c.table_schema and x.table_name = c.table_name and x.table_type = 'BASE TABLE'
     where c.table_schema in ('scm', 'public') and c.data_type = 'text'
       and c.column_name in ('item_code', 'product_code', 'new_item_code', 'trigger_item_code', 'redeemed_item_code', 'sku', 'ac_item_code')
       and not (c.table_schema = 'scm' and c.table_name in ('supplier_material_bindings', 'supplier_binding_price_history'))
  loop
    execute format('update pg_temp.hq_retire r set refs = r.refs + q.n, used_in = r.used_in || %L || '' ''
                      from (select x.%I as code, count(*) as n from %I.%I x where x.%I in (select code from pg_temp.hq_retire) group by 1) q
                     where q.code = r.code', t.s || '.' || t.tb, t.col, t.s, t.tb, t.col);
  end loop;
end $$;
update hq_retire r set refs = r.refs + q.n, used_in = r.used_in || 'scm.mfg_product_photos '
  from (select product_id, count(*) as n from scm.mfg_product_photos group by 1) q
 where q.product_id = r.id;

insert into scm.entity_audit_log (company_id, entity_type, entity_id, action, actor_name_snapshot, field_changes, source, note)
select {CO}, 'mfg_product', r.code, case when r.refs = 0 then 'delete' else 'deactivate' end, 'migration 0142', '[]'::jsonb, 'migration',
       '{TAG}: not on {HK_NOTE} (the quotation replaces the HOOKKA catalogue)'
       || case when r.refs = 0 then '; nothing named it, deleted with its supplier link and cost history'
               else '; still named in ' || trim(r.used_in) || ', so deactivated instead of deleted' end
  from hq_retire r;
delete from scm.supplier_binding_price_history h using hq_retire r
 where r.refs = 0 and h.company_id = {CO} and h.material_kind = 'mfg_product' and h.item_code = r.code;
delete from scm.supplier_material_bindings b using hq_retire r
 where r.refs = 0 and b.company_id = {CO} and b.material_kind = 'mfg_product' and b.item_code = r.code;
delete from scm.mfg_products p using hq_retire r where r.refs = 0 and p.id = r.id;
update scm.mfg_products p set status = 'INACTIVE', updated_at = now() from hq_retire r where r.refs > 0 and p.id = r.id and p.status <> 'INACTIVE';
-- A HOOKKA model left with no SKU and not in the quotation (e.g. 1009, now 1009(A)) existed only for those SKUs.
delete from scm.product_models m
 where m.company_id = {CO} and m.model_code like any(array['BED-HKA-%', 'SOF-HKA-%', 'ACC-HKA-%'])
   and exists (select 1 from hq_retire)
   and not exists (select 1 from hq_models x where x.model_code = m.model_code and x.category = m.category::text)
   and not exists (select 1 from hq_sku n where n.code = m.model_code)
   and not exists (select 1 from scm.mfg_products p where p.model_id = m.id);
-- Each family model lists exactly the sizes / compartments its SKUs now have.
update scm.product_models m set allowed_options = m.allowed_options || x.allowed_options, updated_at = now()
  from hq_models x
 where m.company_id = {CO} and m.model_code = x.model_code and m.category::text = x.category
   and m.allowed_options || x.allowed_options is distinct from m.allowed_options;

do $$ begin
  if exists (select 1 from scm.supplier_material_bindings b join pg_temp.hq_retire r on r.code = b.item_code
              where r.refs = 0 and b.company_id = (select id from public.companies where code = 'HOUZS')) then
    raise exception 'hookka quote: a supplier link still names a deleted SKU';
  end if;
  if exists (select 1 from pg_temp.hq_sku n
              where (select count(*) from scm.supplier_material_bindings b
                      where b.company_id = (select id from public.companies where code = 'HOUZS') and b.material_kind = 'mfg_product'
                        and b.item_code = n.code and b.is_main_supplier) <> 1) then
    raise exception 'hookka quote: a SKU does not have exactly one main supplier link';
  end if;
end $$;

-- 6. HOOKKA's surcharges as HOOKKA's own price book (the PO reads scope supplier:<id> before the master).
insert into scm.maintenance_config_history (id, scope, config, effective_from, notes, company_id)
select 'mch-hookka-300t-20261022', 'supplier:' || s.supplier_id,
       mc.config || (select jsonb_object_agg(k.key,
                        k.value || coalesce((select jsonb_agg(e) from jsonb_array_elements(coalesce(mc.config -> k.key, '[]'::jsonb)) e
                                              where not exists (select 1 from jsonb_array_elements(k.value) x
                                                                 where coalesce(x ->> 'value', x #>> '{{}}') = coalesce(e ->> 'value', e #>> '{{}}'))),
                                             '[]'::jsonb))
                       from jsonb_each({lit(supplier_pools)}::jsonb) k),
       date '{EFF}', '{HK_NOTE}: HOOKKA''s own surcharges (customer override effective 05 May 2026); options it does not quote keep the master value',
       {CO}
  from hq_sup s
  cross join lateral (select coalesce((select m.config from scm.maintenance_config_history m
                                         where m.company_id = {CO} and m.scope = 'master' and m.effective_from <= date '{EFF}'
                                         order by m.effective_from desc, m.created_at desc limit 1), '{{}}'::jsonb) as config) mc
 where s.key = 'HOOKKA' and s.supplier_id is not null and exists (select 1 from hq_sku)
on conflict (id) do nothing;

-- 7. Customer-supplied placeholders: supplier CUSTOMER, one "Customer <sub-category>" SKU per active sub-category and
--    one per group, each linked to CUSTOMER as main supplier at cost 0. A renovation contract never raises a PO for them
--    (engine/backend/src/scm/lib/renovation-po-split.ts isCustomerSupplied).
insert into scm.suppliers (company_id, code, name, status, notes)
select {CO}, 'CUSTOMER', 'Customer', 'ACTIVE', 'customer-supplied items managed by The Conts'
 where exists (select 1 from public.companies where code = 'HOUZS')
   and not exists (select 1 from scm.suppliers s where s.company_id = {CO} and s.code = 'CUSTOMER');

create temp table hq_cus on commit drop as
select 'CUS-' || s.code as code, 'Customer ' || s.label as name, s.group_code as category, s.code as sub_category
  from scm.product_sub_categories s
 where s.active and exists (select 1 from public.companies where code = 'HOUZS')
union all
select g.code, g.name, g.grp, coalesce(
         (select s.code from scm.product_sub_categories s where s.group_code = g.grp and s.code = g.sub and s.active),
         case when g.grp <> 'SERVICE' then (select s.code from scm.product_sub_categories s where s.group_code = g.grp and s.active
                                             order by s.sort_order, s.code limit 1) end)
  from jsonb_to_recordset({lit([dict(grp=a, code=b, name=c, sub=d) for a, b, c, d in CUS_GROUPS])}::jsonb) as g(grp text, code text, name text, sub text)
 where exists (select 1 from public.companies where code = 'HOUZS')
   and exists (select 1 from scm.product_category_groups x where x.code = g.grp);

insert into scm.product_models (company_id, model_code, name, category)
select {CO}, c.code, c.name, c.category::scm.mfg_product_category from hq_cus c
on conflict (company_id, model_code, category) do nothing;
insert into scm.mfg_products (id, company_id, code, name, category, sub_category, uom, status, cost_price_sen, description, model_id)
select 'mfg-' || left(encode(sha256(convert_to(c.code, 'UTF8')), 'hex'), 12), {CO}, c.code, c.name, c.category::scm.mfg_product_category,
       c.sub_category, 'UNIT', 'ACTIVE', 0, '{CUS_DESC}', m.id
  from hq_cus c
  left join scm.product_models m on m.company_id = {CO} and m.model_code = c.code and m.category::text = c.category
 where not exists (select 1 from scm.mfg_products p where p.company_id = {CO} and p.code = c.code)
on conflict do nothing;
insert into scm.supplier_material_bindings
    (company_id, supplier_id, material_kind, item_code, material_name, supplier_sku, unit_price_sen, price_valid_from, is_main_supplier, notes)
select {CO}, s.id, 'mfg_product', c.code, c.name, c.code, 0, date '{EFF}',
       not exists (select 1 from scm.supplier_material_bindings x where x.company_id = {CO} and x.material_kind = 'mfg_product'
                                                                  and x.item_code = c.code and x.is_main_supplier),
       '{CUS_NOTE}'
  from hq_cus c
  join scm.suppliers s on s.company_id = {CO} and s.code = 'CUSTOMER'
 where exists (select 1 from scm.mfg_products p where p.company_id = {CO} and p.code = c.code)
   and not exists (select 1 from scm.supplier_material_bindings b where b.company_id = {CO} and b.material_kind = 'mfg_product'
                                                                  and b.item_code = c.code and b.supplier_id = s.id);

insert into scm.app_config (key, value, description, company_id)
select 'note:{TAG.replace(' ', '-')}',
       json_build_object('tag', '{TAG}', 'at', now(), 'skus', (select count(*) from hq_sku),
                         'priced', (select count(*) from hq_sku where priced),
                         'customer_skus', (select count(*) from hq_cus),
                         'deleted', (select coalesce(array_agg(code order by code), '{{}}') from hq_retire where refs = 0),
                         'deactivated', (select coalesce(array_agg(code order by code), '{{}}') from hq_retire where refs > 0))::text,
       'Catalog note: {HK_NOTE} + AKEMI mattresses (migration 0142)', {CO}
 where exists (select 1 from hq_sku) or exists (select 1 from hq_cus)
on conflict (key) do nothing;

notify pgrst, 'reload schema';
"""
open(MIG, 'w', encoding='utf-8', newline='\n').write(sql)

# ---------------------------------------------------------------- review workbook
HEAD = Font(bold=True); FILL = PatternFill('solid', fgColor='DDE6F0'); WARN = PatternFill('solid', fgColor='FCE4D6')
wb = openpyxl.Workbook()
def sheet(title, head, widths):
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in ws[1]:
        c.font = HEAD; c.fill = FILL; c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = 'A2'
    return ws
def money(ws, cols):
    for col in cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in row:
                c.number_format = '#,##0.00'

rm_ = wb.active; rm_.title = 'Read me'
priced = sum(1 for s in skus if s['priced'])
for line in [
    'HOOKKA quotation to The Conts (customer 300-T, effective 22 Oct 2026, 16 pages) as the HOOKKA catalog, and the AKEMI mattresses '
    'from Houzs Century\'s mattress list of 22 Oct 2026. These are OUR COSTS (what the supplier charges The Conts).',
    f'SKUs after migration 0142: {fam["BEDFRAME"]} bedframes, {fam["SOFA"]} sofa compartments, {fam["ACCESSORY"]} HOOKKA accessories, '
    f'{fam["MATTRESS"]} AKEMI mattresses = {len(skus)}; {priced} with a cost, {len(skus) - priced} "price needed".',
    f'HOOKKA SKUs in the system before (migration 0132 + first import): 163; kept and updated: {sum(1 for s in skus if s["exists"])}; '
    f'removed (not on the quotation): {len(retired)}; created: {sum(1 for s in skus if not s["exists"])}.',
    'Every SKU has ONE main supplier with its cost: picking the SKU on a purchase order fills the supplier and the cost.',
    'Bedframe cost: Price 2 (normal fabric) and Price 1 (cheaper fabric tier). Sofa cost: per seat height (24" ... 35") and fabric tier P2 / P3.',
    'A sofa line with no seat height (a renovation contract line) takes the P2 cost at the smallest quoted seat height (see Questions).',
    'Sheets: Bedframes, Sofas (matrix), Accessories, Mattresses, Surcharges, Changes vs current system, Customer items, Questions.',
    'Source data: docs/catalog-import/build/hookka_300t.json (read from the PDF by build/hookka_parse.py) and build/akemi_mattress.csv.']:
    rm_.append([line])
rm_.column_dimensions['A'].width = 140

ws = sheet('Bedframes', ['SKU code', 'HOOKKA code', 'Description', 'Model', 'Size code', 'Size', 'Price 2 (RM)', 'Price 1 (RM)',
                         'PDF page', 'In system before?'], [30, 26, 50, 26, 9, 14, 12, 12, 9, 14])
for s in skus:
    if s['family'] == 'BEDFRAME':
        ws.append([s['code'], s['supplier_sku'], s['name'], s['base_model'], s['size_code'], s['size_label'], rm(s['quote_p2']),
                   rm(s['quote_p1']), s['page'], 'yes (updated)' if s['exists'] else 'no (new)'])
money(ws, [7, 8]); ws.auto_filter.ref = ws.dimensions

allh = ['24', '26', '28', '30', '32', '35']
ws = sheet('Sofas (matrix)', ['SKU code', 'Model', 'Compartment', 'Tier'] + [f'{h}" (RM)' for h in allh] +
           ['Flat cost used without a seat height (RM)', 'PDF page', 'In system before?'], [26, 8, 12, 6] + [11] * 6 + [16, 9, 14])
for s in skus:
    if s['family'] != 'SOFA':
        continue
    for t in sorted(s['tiers']):
        ws.append([s['code'], s['base_model'].replace('SOF-HKA-', ''), s['code'].split('-')[-1], t] +
                  [rm(s['tiers'][t].get(h)) for h in allh] + [rm(s['unit_price_sen']) if t == 'P2' else None, s['page'],
                   'yes (updated)' if s['exists'] else 'no (new)'])
money(ws, range(5, 12)); ws.auto_filter.ref = ws.dimensions

ws = sheet('Accessories', ['SKU code', 'HOOKKA code', 'Name', 'Cost (RM)', 'PDF page', 'In system before?'], [28, 20, 40, 10, 9, 14])
for s in skus:
    if s['family'] == 'ACCESSORY':
        ws.append([s['code'], s['supplier_sku'], s['name'], rm(s['unit_price_sen']), s['page'], 'yes (updated)' if s['exists'] else 'no (new)'])
money(ws, [4])

ws = sheet('Mattresses', ['SKU code (= supplier code)', 'Name', 'Model (range)', 'Size code', 'Size', 'Cost (RM)', 'Price 1 (RM)', 'Status'],
           [34, 50, 18, 9, 12, 11, 11, 16])
for s in skus:
    if s['family'] == 'MATTRESS':
        ws.append([s['code'], s['name'], s['base_model'], s['size_code'], s['size_label'], rm(s['base_price_sen']), rm(s['price1_sen']),
                   'cost from list' if s['priced'] else 'PRICE NEEDED'])
        if not s['priced']:
            for c in ws[ws.max_row]:
                c.fill = WARN
money(ws, [6, 7]); ws.auto_filter.ref = ws.dimensions

ws = sheet('Surcharges', ['Group', 'List', 'Option', 'HOOKKA surcharge (RM)', 'Products > Maintenance today (RM, from Houzs, mig 0129)',
                          'Compare'], [10, 14, 26, 16, 22, 50])
for r in surch_rows:
    ws.append(r)
    if r[5] == 'DIFFERENT':
        for c in ws[ws.max_row]:
            c.fill = WARN
ws.append([])
ws.append(['What was done: HOOKKA\'s figures are saved as HOOKKA\'s own price book (Maintenance, scope "supplier: HOOKKA", effective '
           '22 Oct 2026). A HOOKKA purchase order adds these surcharges to the cost. Options HOOKKA does not quote keep the master '
           'value. The master (Products > Maintenance) is NOT changed: sales-order costing still reads it.'])
ws.append(['HOOKKA also prints sofa combo prices for 5535 (3 combos, P2): not imported (owner rule of 20 Oct: Conts sets its own combos). '
           'They are on the Questions sheet.'])

ws = sheet('Changes vs current system', ['SKU code', 'Change', 'Before', 'After'], [30, 18, 44, 60])
for s in skus:
    if s['exists']:
        o = OLD[s['code']]
        ws.append([s['code'], 'updated', f'{o["name"]}; cost RM 0.00 (price needed)',
                   f'{s["name"]}; cost RM {rm(s["unit_price_sen"]):,.2f} from the quotation' if s['priced'] else s['name']])
for s in skus:
    if not s['exists']:
        ws.append([s['code'], 'created', '-', f'{s["name"]}; ' + (f'cost RM {rm(s["unit_price_sen"]):,.2f}' if s['priced'] else 'PRICE NEEDED')
                   + f' ({s["supplier"]})'])
for c in retired:
    ws.append([c, 'removed', OLD[c]['name'], 'not on the quotation: deleted (or deactivated if a document names it); written to the audit log'])
    for cell in ws[ws.max_row]:
        cell.fill = WARN
ws.append([])
ws.append(['Surcharges: see the Surcharges sheet (rows marked DIFFERENT).'])
ws.auto_filter.ref = f'A1:D{ws.max_row - 2}'

ws = sheet('Customer items', ['SKU code', 'Name', 'Group', 'Sub-category', 'Unit', 'Supplier', 'Cost (RM)', 'Purchase order'],
           [26, 36, 16, 26, 7, 12, 10, 44])
for sc in SUBCATS:
    ws.append(['CUS-' + sc['code'], 'Customer ' + sc['label'], sc['group_label'], sc['label'], 'UNIT', 'Customer', 0,
               'never raised from a renovation contract'])
labels = {sc['code']: sc['label'] for sc in SUBCATS}
glabels = {sc['group_code']: sc['group_label'] for sc in SUBCATS}
for grp, code, name, sub in CUS_GROUPS:
    ws.append([code, name, glabels.get(grp, grp), (labels.get(sub, '') + ' (group item)').strip(), 'UNIT', 'Customer', 0,
               'never raised from a renovation contract'])
money(ws, [7])
ws.append([])
ws.append([f'{len(SUBCATS)} sub-categories on a database built from the migrations (production has 55 after 0141): the migration '
           'reads the sub-categories when it runs, so every active one on production gets its Customer SKU.'])
ws.append(['Meaning: the customer orders / handles this item; The Conts manages it on the project. Supplier CUSTOMER, cost RM 0.00 on purpose '
           '(not "price needed"). A renovation contract line with a Customer SKU is skipped when purchase orders are raised.'])

ws = sheet('Questions', ['#', 'Question for the owner', 'What the system does until answered'], [4, 100, 70])
QS = [
    ('HOOKKA\'s section header says SOFA (209), and its per-model headers add up to 155, but the quotation prints 145 model x '
     'compartment rows (226 rows counting P2 and P3). 5545 / 5546 / 5547 say 14 and 5548 / 5549 say 15 but print 12-13. '
     'Are some compartments missing from the quotation?', 'Only the printed rows are SKUs: 194 sofa SKUs (145 rows, "(LHF/RHF)" split into a left and a right SKU).'),
    ('5536 CNR, P2, 30": RM 8,258.00 (24" is RM 825, 28" RM 900). Typo for RM 825.80 or RM 858?', 'Stored as printed (RM 8,258.00). Correct it in SKU Master or ask HOOKKA.'),
    ('5549 L(RHF), P2, 26": RM 100.00 (L(LHF) at 26" is RM 1,000). Typo for RM 1,000?', 'Stored as printed (RM 100.00).'),
    ('5539 1B / 2B: P3 is cheaper than P2 and has no 26" price; 5539 1B P2 24" RM 800 vs P3 RM 700. Swapped tiers?', 'Stored as printed.'),
    ('L (LHF/RHF) of 5530 / 5531 / 5535 / 5536: 35" is cheaper than 32" (e.g. 5530 P2 RM 1,200 at 32", RM 1,010 at 35").', 'Stored as printed.'),
    ('Odd cents: 1003(A)(HF)(W)-(SS) RM 582.99, DIVAN-(SS) RM 319.02, 5548 L(RHF) 35" RM 1,249.99.', 'Stored as printed.'),
    ('5545, 5546, 5547, 5548, 5549 and 6112 have only P2 (no P3). A P3 fabric on those takes the P2 cost.', 'P2 cost used for every fabric.'),
    ('A sofa line without a seat height (renovation contract lines have none) needs one flat cost. Which seat height should it be?',
     'P2 at the smallest quoted seat height (24", or 30" for 6112). A line with a seat height picked uses the exact matrix cell.'),
    ('Sofa combos for 5535 (P2, 24"-35": RM 1,584-1,880 / 2,100-2,500 / 2,600-3,300) are HOOKKA prices for buying the set. '
     'Set them as HOOKKA combo costs?', 'Not imported; each piece is costed on its own.'),
    ('SQUARE PILLOW SP: description says 26"x26", size column says 16"x16" (RM 55 vs RM 25 for 16"x16").', 'Imported as RM 55.00, name kept.'),
    ('Compartments CSL (5536, 5540, 5543, 5545-5549) and SQCNR (5543) are not in the sofa builder\'s compartment library (no picture).',
     'SKUs created; the sofa builder shows them without a drawing.'),
    ('Customer items per group sit under one sub-category each (Works > Protection, Built-in > Shelving, Furniture > Loose storage, '
     'Lighting > Downlight, Appliances > TV, Soft furnishing > Decoration; Service needs none). The engine needs a sub-category for every '
     'group but Service. Better homes for them?', 'As listed on the Customer items sheet.'),
    ('HOOKKA lists sofa seat sizes 44, 24" x 37" and 12 that no quoted price uses.', 'Not added to the sizes list.'),
    ('Surcharges that differ from Products > Maintenance (copied from Houzs): see the Surcharges sheet. Should the master take HOOKKA\'s values?',
     'HOOKKA purchase orders use HOOKKA\'s values; sales-order costing keeps the master values.'),
    (f'{len(retired)} HOOKKA SKUs of the old catalogue are not on the quotation (e.g. 1003-(SK) and 1003-(220X220): the quotation has those sizes only '
     'under 1003(A); the 1009 family is now 1009(A); DIVAN sizes are coded by HOOKKA as DIVAN-(153) etc.).', 'Removed (see Changes).'),
    (f'AKEMI: {sum(1 for s in skus if s["family"] == "MATTRESS" and not s["priced"])} rows have no price: every CUSTOM (SP) size, '
     'plus HS SOLACE 200x200 and all 5 KORI IYASHI sizes.', 'Created and linked to Houzs Century with cost 0, "PRICE NEEDED".'),
]
for i, (q, a) in enumerate(QS, 1):
    ws.append([i, q, a])
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(os.path.join(DOCS, 'hookka-300t-review.xlsx'))
print(fam, 'priced', priced, 'retired', len(retired), 'models', len(models))
