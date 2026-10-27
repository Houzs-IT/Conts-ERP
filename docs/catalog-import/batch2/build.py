"""Build the second SKU batch (26 Oct 2026): the balance SKU file and the codes opened from the AutoCount SOs.
Usage: python -I -X utf8 docs/catalog-import/batch2/build.py <balance.xlsx> <New-SKU-codes.xlsx> <out_dir>
Writes <out_dir>/batch2-data.json, batch2-a-balance.sql, batch2-b-so-codes.sql."""
import sys, json, hashlib, collections, os
import openpyxl

bal_x, new_x, out = sys.argv[1:4]
TAG = 'sku-batch2 2026-10-26'
TODAY = '2026-10-26'

# Balance-file supplier name -> Conts supplier (code, name). The last four are new.
SUP = {
    'SHOPEE': ('400-S002', 'Shopee'), 'IKEA': ('400-I001', 'IKEA'), 'ZOE': ('400-Z001', 'ZOE HOME SDN BHD'),
    'UJU': ('400-U001', 'UJU HOME'), 'VIP HOSPITALITY': ('400-V001', 'VIP HOSPITALITY SUPPLIES SDN BHD'),
    'HLK': ('400-0002', 'HLK (CHAIN-STORE) SDN BHD'), 'ONE LIVING': ('400-O003', 'ONE LIVING SDN BHD'),
    'SINCERE ZONE': ('400-S001', 'SINCERE ZONE ELECTRONICS (M) SDN BHD'), 'CARRESS': ('400-C004', 'CARRESS SDN BHD'),
    'HOUZS CENTURE': ('400-H003', 'HOUZS CENTURY SDN BHD'),
    'HOMES HARMONY': ('400-H011', 'HOMES HARMONY'), 'KAISON': ('400-K007', 'KAISON'),
    'DIY': ('400-D005', 'DIY'), 'OFFICE OWN DO': ('400-O004', 'OFFICE OWN DO'),
}
NEW_SUP = {'HOMES HARMONY', 'KAISON', 'DIY', 'OFFICE OWN DO'}

# Sub-categories this batch adds at run time (settings-editable groups, like WALL_CABINET / TRACK_LIGHT before).
NEW_SUBS = [
    ('APPLIANCE', 'SMALL_APPLIANCE', 'Small appliance', 70),
    ('SOFT', 'TOWEL_BATH', 'Towel & bath', 70),
    ('SERVICE', 'DESIGN', 'Design', 80),
    ('SERVICE', 'SIGNBOARD', 'Signboard', 90),
    ('SERVICE', 'DEFECT_CHECK', 'Defect checking', 100),
    # Already on production (added in settings); listed so a fresh database (local check, staging) has them too.
    ('BUILTIN', 'WALL_CABINET', 'Wall Cabinet', 80),
    ('LIGHTING', 'TRACK_LIGHT', 'Track light', 200),
]

def place(code, name, sub):
    """Balance-file (group, sub-category) -> our (category, sub_category), as proposed to the owner on 24 Oct."""
    if sub == 'Appliances': return 'APPLIANCE', 'SMALL_APPLIANCE'
    if sub == 'Bedding': return ('SOFT', 'TOWEL_BATH') if code in ('LIN-0001', 'LIN-0014') else ('SOFT', 'BED_LINEN')
    if sub == 'Carpet': return 'SOFT', 'CARPET'
    if code == 'DEC-0009': return 'SOFT', 'DIFFUSER'
    if sub in ('Photo', 'Small Deco'): return 'SOFT', 'DECORATION'
    return 'AIRBNB', {'Kitchen': 'KITCHENWARE', 'Bathroom': 'BATH_CLEANING', 'Fixture': 'FIXTURE',
                      'Consume': 'CONSUMABLES', 'Signage': 'SIGNAGE'}[sub]

UOM = {'PCS': 'PCS', 'PACK': 'PKT', 'JOB': 'LOT'}

wb = openpyxl.load_workbook(bal_x, data_only=True)
master = [r for r in list(wb['ERP SKU MASTER'].iter_rows(values_only=True))[1:] if r[0]]
mapping = [r for r in list(wb['SKU SUPPLIER MAPPING'].iter_rows(values_only=True))[1:] if r[0]]
services = [r for r in list(wb['SERVICES SEPARATE'].iter_rows(values_only=True))[1:] if r[0]]
DROP = {'AIB-0037'}  # same item, name and price as AIB-0036 (owner told 24 Oct)

bal = []
for code, name, grp, sub, uom, _sup, cost, _st in master:
    if code in DROP: continue
    cat, sc = place(code, name, sub)
    links = [(s, prim == 'Yes') for (c, _n, _g, _s, s, prim, _cst, _note) in mapping if c == code]
    assert sum(p for _, p in links) == 1, code
    bal.append(dict(code=code, name=name, category=cat, sub_category=sc, uom=UOM[uom], cost_rm=cost,
                    links=[dict(supplier=SUP[s][0], supplier_name=SUP[s][1], main=p) for s, p in links],
                    source=f'balance file: {grp} / {sub}'))
for code, name, grp, sub, uom, _sup, cost, _st in services:
    bal.append(dict(code=code, name=name, category='SERVICE', sub_category='AIRBNB_SETUP', uom=UOM[uom],
                    cost_rm=cost if isinstance(cost, (int, float)) else None, links=[], source=f'balance file: {grp} / {sub}'))

nw = openpyxl.load_workbook(new_x, data_only=True).active
so = []
SERVICE_SUB = {'RN-DESIGN': 'DESIGN', 'RN-SIGNAGE': 'SIGNBOARD', 'RN-DEFECT': 'DEFECT_CHECK'}
for r in list(nw.iter_rows(values_only=True))[1:]:
    if not r[1]: continue
    _no, code, name, cat, sub, uom = r[:6]
    if cat == 'SERVICE': sub = SERVICE_SUB[code]
    so.append(dict(code=code, name=name, category=cat, sub_category=sub, uom=uom, cost_rm=None, links=[],
                   source='outstanding AutoCount SOs, New-SKU-codes.xlsx'))

allc = [k['code'] for k in bal + so]
dup = [c for c, n in collections.Counter(allc).items() if n > 1]
assert not dup, dup
names = collections.Counter(k['name'].strip().upper() for k in bal + so)
assert not [n for n, c in names.items() if c > 1], [n for n, c in names.items() if c > 1]

os.makedirs(out, exist_ok=True)
json.dump(dict(built=TODAY, tag=TAG, new_sub_categories=NEW_SUBS,
               new_suppliers=[dict(key=k, code=SUP[k][0], name=SUP[k][1]) for k in sorted(NEW_SUP)],
               balance=bal, so_codes=so), open(os.path.join(out, 'batch2-data.json'), 'w', encoding='utf-8', newline='\n'),
          indent=1, ensure_ascii=False)

def lit(v): return 'null' if v is None else "'" + str(v).replace("'", "''") + "'"
def sid(code): return 'mfg-' + hashlib.sha256(f'sku-batch2|{code}'.encode()).hexdigest()[:12]
COMPANY = "(select id from public.companies where code = 'HOUZS')"
Q = '$b2$'

def part(label, skus, with_new_suppliers):
    rows = [dict(id=sid(k['code']), code=k['code'], name=k['name'], category=k['category'], sub_category=k['sub_category'],
                 uom=k['uom'], cost=round((k['cost_rm'] or 0) * 100)) for k in skus]
    links = [dict(code=k['code'], name=k['name'], sup_code=l['supplier'], main=l['main'],
                  cost=round((k['cost_rm'] or 0) * 100) if l['main'] else 0,
                  priced=bool(l['main'] and k['cost_rm']))
             for k in skus for l in k['links']]
    codes = 'array[' + ','.join(lit(k['code']) for k in skus) + ']::text[]'
    subs = ',\n  '.join(f"({lit(g)}, {lit(c)}, {lit(l)}, {s}, '{{}}'::text[])" for g, c, l, s in NEW_SUBS)
    sups = ''
    if with_new_suppliers:
        vals = ',\n  '.join(f"({lit(SUP[k][0])}, {lit(SUP[k][1])})" for k in sorted(NEW_SUP))
        sups = f"""
-- Suppliers named in the balance file that Conts did not have yet.
insert into scm.suppliers (company_id, code, name, notes, status)
select {COMPANY}, x.code, x.name, {lit('Created by ' + TAG + ' from the balance SKU file.')}, 'ACTIVE'
  from (values
  {vals}) as x(code, name)
 where not exists (select 1 from scm.suppliers s where s.company_id = {COMPANY} and (upper(s.code) = upper(x.code) or upper(trim(s.name)) = upper(x.name)));
"""
    return f"""-- {TAG}: {label} ({len(skus)} SKUs). Built by docs/catalog-import/batch2/build.py from batch2-data.json.
-- One transaction; safe to run again (everything keyed on SKU code / supplier code). Needs migrations 0162 and 0163.
begin;

do $$ begin
  if (select count(*) from public.companies where code = 'HOUZS') <> 1 then raise exception 'expected exactly one HOUZS company'; end if;
  if not exists (select 1 from scm.product_category_groups where code = 'AIRBNB') then raise exception 'apply migration 0163 first'; end if;
end $$;

insert into scm.product_sub_categories (group_code, code, label, sort_order, legacy_categories) values
  {subs}
on conflict do nothing;
{sups}
create temp table b2_rows on commit drop as
select * from jsonb_to_recordset({Q}{json.dumps(rows, ensure_ascii=False)}{Q}::jsonb)
  as x(id text, code text, name text, category text, sub_category text, uom text, cost integer);

-- Every SKU belongs to a model of its own (model_code = SKU code), as the API's create paths do.
insert into scm.product_models (company_id, model_code, name, category)
select {COMPANY}, r.code, r.name, r.category::scm.mfg_product_category from b2_rows r
on conflict (company_id, model_code, category) do nothing;

-- Active straight away: photos are optional since migration 0130.
insert into scm.mfg_products (id, company_id, code, name, category, sub_category, uom, status, cost_price_sen, model_id)
select r.id, {COMPANY}, r.code, r.name, r.category::scm.mfg_product_category, r.sub_category, r.uom, 'ACTIVE', r.cost, m.id
  from b2_rows r
  left join scm.product_models m on m.company_id = {COMPANY} and m.model_code = r.code and m.category::text = r.category
on conflict (company_id, code) do nothing;

create temp table b2_links on commit drop as
select x.*, (select s.id from scm.suppliers s where s.company_id = {COMPANY} and upper(s.code) = upper(x.sup_code)) as supplier_id
  from jsonb_to_recordset({Q}{json.dumps(links, ensure_ascii=False)}{Q}::jsonb)
  as x(code text, name text, sup_code text, main boolean, cost integer, priced boolean);

do $$ begin
  if exists (select 1 from b2_links where supplier_id is null) then raise exception 'a supplier code in this batch does not exist'; end if;
end $$;

-- One link per supplier and SKU; the balance file's primary supplier is main (cost = its Main Cost), the others
-- are alternatives with no price yet (TBC in the file).
insert into scm.supplier_material_bindings
    (company_id, supplier_id, material_kind, item_code, material_name, supplier_sku, unit_price_sen, price_valid_from, is_main_supplier, notes)
select {COMPANY}, l.supplier_id, 'mfg_product', l.code, l.name, '', l.cost, case when l.priced then date {lit(TODAY)} end,
       l.main and not exists (select 1 from scm.supplier_material_bindings b where b.company_id = {COMPANY} and b.material_kind = 'mfg_product'
                                                                            and b.item_code = l.code and b.is_main_supplier),
       {lit(TAG + ': ')} || case when l.main then 'main supplier, cost from the balance file' else 'alternative supplier, price TBC' end
  from b2_links l
 where not exists (select 1 from scm.supplier_material_bindings b where b.company_id = {COMPANY} and b.material_kind = 'mfg_product'
                                                                   and b.item_code = l.code and b.supplier_id = l.supplier_id);

insert into scm.supplier_binding_price_history
    (company_id, supplier_id, material_kind, item_code, unit_price_sen, is_main_supplier, effective_from, notes, created_by, applied_at)
select {COMPANY}, l.supplier_id, 'mfg_product', l.code, l.cost, true, date {lit(TODAY)}, {lit(TAG + ': balance file Main Cost')}, {lit(TAG)}, now()
  from b2_links l
 where l.priced
   and not exists (select 1 from scm.supplier_binding_price_history h where h.company_id = {COMPANY} and h.supplier_id = l.supplier_id
                     and h.material_kind = 'mfg_product' and h.item_code = l.code and h.effective_from = date {lit(TODAY)});

insert into scm.app_config (key, value, description, company_id)
values ({lit('note:sku-batch2-2026-10-26:' + label.split(' ')[0])}, {lit(json.dumps(dict(tag=TAG, part=label, skus=len(skus))))},
        {lit('Import note: ' + TAG + ', ' + label)}, {COMPANY})
on conflict (key) do nothing;

commit;

select
  (select count(*)::int from scm.mfg_products where company_id = {COMPANY} and code = any({codes})) as skus,
  {len(skus)} as skus_expected,
  (select count(*)::int from scm.supplier_material_bindings where company_id = {COMPANY} and material_kind = 'mfg_product' and item_code = any({codes})) as supplier_links,
  {len(links)} as links_expected,
  (select count(*)::int from scm.mfg_products where company_id = {COMPANY}) as all_skus;
"""

open(os.path.join(out, 'batch2-a-balance.sql'), 'w', encoding='utf-8', newline='\n').write(part('balance SKU file', bal, True))
open(os.path.join(out, 'batch2-b-so-codes.sql'), 'w', encoding='utf-8', newline='\n').write(part('SO codes', so, False))
print('balance', len(bal), 'links', sum(len(k['links']) for k in bal), '| so', len(so))
print(collections.Counter((k['category'], k['sub_category']) for k in bal))
