"""Outstanding SO lines -> a catalog SKU each: an existing one, one from the balance SKU file, or a proposed new one.

Totals, discounts, room headings and size lines are left out. Every rule is (target, regex), first match wins."""
import json, re, collections, statistics, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

SRC = r"C:\Users\User\Downloads\outstanding SO.xlsx"
BAL = r"C:\Users\User\Downloads\THE_CONTS_ERP_SKU_COSTS_UPDATED.xlsx"
HERE = r"C:/Users/User/AppData/Local/Temp/claude/C--Users-User-Desktop-Conts-ERP/9f6e16ec-e81b-495c-be84-66e8a1620520/scratchpad"
OUT = None

cat = {c[0]: c for c in json.load(open(HERE + "/catalog.json", encoding="utf-8"))}
bwb = openpyxl.load_workbook(BAL, data_only=True, read_only=True)
bal = {}
for sh in ("ERP SKU MASTER", "SERVICES SEPARATE"):
    for r in list(bwb[sh].iter_rows(values_only=True))[1:]:
        if r[0] and r[0] != "AIB-0037":
            bal[r[0]] = (r[1], r[3], r[4], r[6])

# Proposed new SKUs: code -> (name, group, sub-category, uom, why the unit)
NEW = {
    # built-in carpentry made by Conts's own carpenter (supplier code CTS)
    "BIN-CTS-BASE": ("Kitchen base cabinet", "BUILTIN", "KITCHEN_CABINET", "FT", ""),
    "BIN-CTS-WALL": ("Kitchen wall cabinet", "BUILTIN", "WALL_CABINET", "FT", ""),
    "BIN-CTS-TALL": ("Tall cabinet", "BUILTIN", "KITCHEN_CABINET", "UNIT", ""),
    "BIN-CTS-TALL-GLASS": ("Tall cabinet with 4G glass door", "BUILTIN", "KITCHEN_CABINET", "UNIT", ""),
    "BIN-CTS-FRIDGE": ("Fridge cabinet", "BUILTIN", "KITCHEN_CABINET", "UNIT", ""),
    "BIN-CTS-DRAWER-1": ("Soft close drawer, 1 layer", "BUILTIN", "KITCHEN_CABINET", "UNIT", ""),
    "BIN-CTS-DRAWER-3": ("Soft close drawer, 3 layers", "BUILTIN", "KITCHEN_CABINET", "UNIT", ""),
    "BIN-CTS-DRAWER-UM1": ("Undermount soft close drawer, 1 layer", "BUILTIN", "KITCHEN_CABINET", "UNIT", ""),
    "BIN-CTS-DRAWER-UM3": ("Undermount soft close drawer, 3 layers", "BUILTIN", "KITCHEN_CABINET", "UNIT", ""),
    "BIN-CTS-DISHRACK": ("Stainless steel dish rack 60cm", "BUILTIN", "KITCHEN_CABINET", "UNIT", ""),
    "BIN-CTS-PULLOUT": ("Pull-out tray / basket", "BUILTIN", "KITCHEN_CABINET", "UNIT", ""),
    "BIN-CTS-BACKPANEL": ("Back panel", "BUILTIN", "KITCHEN_CABINET", "UNIT", ""),
    "BIN-CTS-FOAMBOARD": ("Waterproof foam board", "BUILTIN", "KITCHEN_CABINET", "FT", ""),
    "BIN-CTS-ALUCASING": ("Aluminium LED casing", "BUILTIN", "KITCHEN_CABINET", "FT", ""),
    "BIN-CTS-EP": ("Melamine end panel (EP)", "BUILTIN", "KITCHEN_CABINET", "UNIT", ""),
    "BIN-CTS-QUARTZ": ("Quartz stone top", "BUILTIN", "COUNTERTOP", "FT", "colour is the LB stone SKU when known"),
    "BIN-CTS-SINTERED": ("Sintered stone top", "BUILTIN", "COUNTERTOP", "FT", ""),
    "BIN-CTS-WARDROBE": ("Swing door wardrobe", "BUILTIN", "WARDROBE", "FT", ""),
    "BIN-CTS-WARDROBE-3D": ("3 door soft close wardrobe (premium)", "BUILTIN", "WARDROBE", "UNIT", ""),
    "BIN-CTS-SHOE": ("Shoe cabinet", "BUILTIN", "SHOE_CABINET", "FT", ""),
    "BIN-CTS-TV": ("TV console (sitting or hanging)", "BUILTIN", "TV_CABINET", "FT", ""),
    "BIN-CTS-DRESSING": ("Built-in dressing table", "BUILTIN", "STUDY_TABLE", "FT", ""),
    "BIN-CTS-STUDY": ("Built-in study table / working desk", "BUILTIN", "STUDY_TABLE", "FT", ""),
    "BIN-CTS-SHELF": ("Wall / floating shelf", "BUILTIN", "SHELVING", "FT", ""),
    "BIN-CTS-HANGROD": ("Hanging rod", "BUILTIN", "WARDROBE", "FT", ""),
    "BIN-CTS-BEDHEAD": ("Bed head panel", "BUILTIN", "SHELVING", "UNIT", ""),
    "BIN-CTS-CABINET": ("Other built-in cabinet (storage, multipurpose)", "BUILTIN", "SHELVING", "FT", ""),
    "BIN-CTS-CURVE": ("Curve design add-on", "BUILTIN", "SHELVING", "UNIT", ""),
    # electrical labour, per point
    "ELC-CTS-LIGHT-PIPE": ("New lighting point, looping (pipe work)", "WORKS", "ELECTRICAL", "POINT", ""),
    "ELC-CTS-LIGHT-CONCEAL": ("New lighting point (conceal work)", "WORKS", "ELECTRICAL", "POINT", ""),
    "ELC-CTS-SOCKET-1": ("New socket, first point", "WORKS", "ELECTRICAL", "POINT", ""),
    "ELC-CTS-SOCKET-2": ("New socket, each subsequent point", "WORKS", "ELECTRICAL", "POINT", ""),
    "ELC-CTS-SWITCH": ("Supply and install 1-gang switch", "WORKS", "ELECTRICAL", "PCS", ""),
    "ELC-CTS-FANHOOK": ("Supply and install fan hook", "WORKS", "ELECTRICAL", "PCS", ""),
    "ELC-CTS-DOORBELL": ("Supply and install door bell", "WORKS", "ELECTRICAL", "PCS", ""),
    "ELC-CTS-INST-FAN": ("Install ceiling fan (owner supply)", "WORKS", "ELECTRICAL", "UNIT", ""),
    "ELC-CTS-INST-LIGHT": ("Install light (owner supply)", "WORKS", "ELECTRICAL", "POINT", ""),
    "ELC-CTS-INST-HEATER": ("Install water heater", "WORKS", "ELECTRICAL", "UNIT", ""),
    "ELC-CTS-INST-HOODHOB": ("Install hood and hob", "WORKS", "ELECTRICAL", "SET", ""),
    # lighting and fans (brand to fill in; GEN = no fixed supplier yet)
    "FAN-GEN-5B42-R": ("Ceiling fan 5 blade 42 inch with remote", "LIGHTING", "CEILING_FAN", "UNIT", ""),
    "FAN-GEN-5B56-R": ("Ceiling fan 5 blade 56 inch with remote", "LIGHTING", "CEILING_FAN", "UNIT", ""),
    "FAN-GEN-5B56-RL": ("Ceiling fan 5 blade 56 inch with remote and light", "LIGHTING", "CEILING_FAN", "UNIT", ""),
    "FAN-GEN-CEILING": ("Ceiling fan with remote (size to confirm)", "LIGHTING", "CEILING_FAN", "UNIT", ""),
    "FAN-FNZ-FAN": ("FANZO ceiling fan", "LIGHTING", "CEILING_FAN", "UNIT", ""),
    "FAN-FNZ-FANLIGHT": ("FANZO ceiling fan with light", "LIGHTING", "CEILING_FAN", "UNIT", ""),
    "FAN-GEN-WALL": ("Wall fan", "LIGHTING", "CEILING_FAN", "UNIT", ""),
    "DLT-GEN-SURF8-CW": ("Surface downlight 8 inch cool white", "LIGHTING", "DOWNLIGHT", "PCS", ""),
    "DLT-GEN-RECESS": ("Recessed downlight", "LIGHTING", "DOWNLIGHT", "PCS", ""),
    "DLT-GEN-LINEAR600": ("Linear light 600mm", "LIGHTING", "DOWNLIGHT", "PCS", ""),
    "LGT-GEN-T5": ("T5 LED batten", "LIGHTING", "LED_STRIP", "PCS", ""),
    "PND-GEN-DINING": ("Dining hanging light", "LIGHTING", "PENDANT", "PCS", ""),
    "PND-GEN-WALL": ("Wall light", "LIGHTING", "PENDANT", "PCS", ""),
    "EYB-GEN-EYEBALL": ("Eyeball light", "LIGHTING", "EYEBALL", "PCS", ""),
    "TRK-GEN-MAGNETIC": ("Recessed magnetic track", "LIGHTING", "TRACK_LIGHT", "M", ""),
    "TRK-GEN-1M-SET": ("Track light 1m with 1+3 heads", "LIGHTING", "TRACK_LIGHT", "SET", ""),
    # appliances
    "TV-HIS-43": ("Hisense smart TV 43 inch", "APPLIANCE", "TV", "UNIT", ""),
    "TV-HIS-50": ("Hisense smart TV 50 inch", "APPLIANCE", "TV", "UNIT", ""),
    "FRG-HIS-420L": ("Hisense fridge 420L 2 door", "APPLIANCE", "FRIDGE", "UNIT", ""),
    "FRG-GEN-110L": ("Fridge 110L", "APPLIANCE", "FRIDGE", "UNIT", ""),
    "FRG-GEN-200L": ("Fridge 200L", "APPLIANCE", "FRIDGE", "UNIT", ""),
    "FRG-GEN-280L": ("Fridge 280L", "APPLIANCE", "FRIDGE", "UNIT", ""),
    "WSH-MID-2IN1": ("Midea washer dryer 2 in 1, 7.5kg", "APPLIANCE", "WASHER", "UNIT", ""),
    "ACU-HIS-1HP": ("Hisense aircon 1HP (incl. 10ft copper)", "APPLIANCE", "AIRCON_UNIT", "UNIT", ""),
    "ACU-HIS-15HP": ("Hisense aircon 1.5HP (incl. 10ft copper)", "APPLIANCE", "AIRCON_UNIT", "UNIT", ""),
    "HTR-GEN-HEATER": ("Water heater", "APPLIANCE", "WATER_HEATER", "UNIT", ""),
    "HOB-HET-2IN1": ("Hetch 2-in-1 induction and ceramic cooker 3000W", "APPLIANCE", "HOOD_HOB", "UNIT", ""),
    "HOD-ELB-M7598": ("Elba slim hood ESH-M7598", "APPLIANCE", "HOOD_HOB", "UNIT", ""),
    # loose furniture
    "FUR-MUJ-SIDETABLE": ("MUJI side table", "FURNITURE", "TABLES", "PCS", ""),
    "FUR-MUJ-COFFEESET": ("MUJI coffee table set", "FURNITURE", "TABLES", "SET", ""),
    "FUR-MUJ-DINTABLE": ("MUJI dining table 1500mm", "FURNITURE", "DINING", "PCS", ""),
    "FUR-MUJ-DINCHAIR": ("MUJI dining chair", "FURNITURE", "DINING", "PCS", ""),
    "FUR-MUJ-DINBENCH": ("MUJI dining bench", "FURNITURE", "DINING", "PCS", ""),
    "FUR-MUJ-TVCAB": ("MUJI TV cabinet (Malmo 1.8m)", "FURNITURE", "LOOSE_STORAGE", "PCS", ""),
    "FUR-MUJ-WARDROBE3D": ("MUJI 3 door soft close wardrobe", "FURNITURE", "LOOSE_STORAGE", "PCS", ""),
    "FUR-MUJ-WARDROBEOPEN": ("MUJI open shelf wardrobe", "FURNITURE", "LOOSE_STORAGE", "PCS", ""),
    "FUR-GEN-DESKTV": ("Working desk + TV cabinet table", "FURNITURE", "TABLES", "PCS", ""),
    "FUR-GEN-DRESSING": ("Designer dressing table", "FURNITURE", "TABLES", "PCS", ""),
    "FUR-GEN-SIDETABLE": ("Side / bedside table", "FURNITURE", "TABLES", "PCS", ""),
    "FUR-GEN-COFFEE": ("Coffee table", "FURNITURE", "TABLES", "PCS", ""),
    "FUR-GEN-NEATH150": ("NEATH dining table 1.5m", "FURNITURE", "DINING", "PCS", ""),
    "FUR-GEN-DINSET2": ("Dining table + 2 chairs", "FURNITURE", "DINING", "SET", ""),
    "FUR-GEN-DINSET6": ("Dining set 1 table + 6 chairs", "FURNITURE", "DINING", "SET", ""),
    "FUR-GEN-CHAIR": ("Dining chair", "FURNITURE", "DINING", "PCS", ""),
    "FUR-GEN-WORKCHAIR": ("Working / study chair", "FURNITURE", "ACCESSORY", "PCS", ""),
    "FUR-GEN-POUF": ("Designer nordic pouf", "FURNITURE", "ACCESSORY", "PCS", ""),
    "SOF-GEN-PREM25": ("Premium sofa 2.5 seater", "FURNITURE", "SOFA", "UNIT", ""),
    "SOF-GEN-3S": ("SOFA 3S", "FURNITURE", "SOFA", "UNIT", ""),
    "AKEMI GALLENT MATT (Q)": ("AKEMI Hotel Premium Gallent mattress, queen", "FURNITURE", "MATTRESS", "PCS", ""),
    "AKEMI GALLENT MATT (K)": ("AKEMI Hotel Premium Gallent mattress, king", "FURNITURE", "MATTRESS", "PCS", ""),
    "CARRES HOTEL PREMIUM L MATT (S)": ("CARRES Hotel Premium Model L mattress, single", "FURNITURE", "MATTRESS", "PCS", ""),
    # soft furnishing
    "CUR-MUJ-SINGLE": ("MUJI curtain, single layer", "SOFT", "CURTAIN_BLINDS", "SET", "per window, package item"),
    "CUR-MUJ-DOUBLE": ("MUJI curtain, double layer", "SOFT", "CURTAIN_BLINDS", "SET", "per window, package item"),
    "CUR-GEN-SINGLE": ("Curtain single-ply SG pleat", "SOFT", "CURTAIN_BLINDS", "FT", ""),
    "CUR-GEN-DOUBLE": ("Curtain double-ply SG pleat", "SOFT", "CURTAIN_BLINDS", "FT", ""),
    "BLD-GEN-BLINDS": ("Blinds", "SOFT", "CURTAIN_BLINDS", "UNIT", ""),
    "MIR-GEN-ROUND600": ("Round mirror 600mm", "SOFT", "DECORATION", "PCS", ""),
    "MIR-CTS-CUSTOM": ("Custom mirror", "SOFT", "DECORATION", "UNIT", ""),
    # finer spec codes (owner, 26 Oct 2026: as fine as the sub-category's own SKUs go)
    "FUR-GEN-CHAIR-ZOLA": ("ZOLA dining chair", "FURNITURE", "DINING", "PCS", ""),
    "FUR-GEN-CHAIR-NANZO": ("NANZO chair", "FURNITURE", "DINING", "PCS", ""),
    "FUR-GEN-CHAIR-NAIDA": ("NAIDA chair", "FURNITURE", "DINING", "PCS", ""),
    "FUR-GEN-CHAIR-MARLO": ("MARLO chair", "FURNITURE", "DINING", "PCS", ""),
    "FUR-GEN-CHAIR-MACY": ("MACY chair", "FURNITURE", "DINING", "PCS", ""),
    "FUR-GEN-CHAIR-CARTER": ("CARTER dining chair", "FURNITURE", "DINING", "PCS", ""),
    "FUR-GEN-DIN-GERALD180": ("GERALD dining table 1800mm", "FURNITURE", "DINING", "PCS", ""),
    "FUR-GEN-DIN-AIMON": ("AIMON dining table", "FURNITURE", "DINING", "PCS", ""),
    "FUR-GEN-DIN-ROUND180": ("Round dining table 1800mm", "FURNITURE", "DINING", "PCS", ""),
    "FUR-GEN-COF-SANDRA2": ("SANDRA coffee table, set of 2", "FURNITURE", "TABLES", "SET", ""),
    "FUR-GEN-COF-PANRA2": ("PANRA coffee table, set of 2", "FURNITURE", "TABLES", "SET", ""),
    "FUR-GEN-COF-OZZY80": ("OZZY round coffee table 800mm", "FURNITURE", "TABLES", "PCS", ""),
    "FUR-GEN-COF-HIACE": ("HIACE coffee table", "FURNITURE", "TABLES", "PCS", ""),
    "FUR-GEN-SIDE-HIRADO": ("HIRADO side table", "FURNITURE", "TABLES", "PCS", ""),
    "FUR-GEN-BEDSIDE": ("Bedside table", "FURNITURE", "TABLES", "PCS", ""),
    "SOF-GEN-NEXON-3S": ("NEXON sofa 3 seater", "FURNITURE", "SOFA", "UNIT", ""),
    "SOF-GEN-BLATT-3S": ("BLATT sofa 3 seater", "FURNITURE", "SOFA", "UNIT", ""),
    "SOF-HKA-6112-3S": ("SOFA 6112 3S", "FURNITURE", "SOFA", "UNIT", "HOOKKA 6112 has only 2S in the catalog"),
    "SOF-GEN-3S-PUSHBACK": ("Sofa 3 seater push back", "FURNITURE", "SOFA", "UNIT", ""),
    "SOF-GEN-BED": ("Sofa bed", "FURNITURE", "SOFA", "UNIT", ""),
    "SOF-GEN-1S": ("Sofa 1 seater", "FURNITURE", "SOFA", "UNIT", ""),
    "DR MACIO MATT (S)": ("DR MACIO mattress 12 inch, single", "FURNITURE", "MATTRESS", "PCS", ""),
    "DOCTOR K MATT (Q)": ("DOCTOR K mattress, queen", "FURNITURE", "MATTRESS", "PCS", ""),
    "DOCTOR K MATT (SS)": ("DOCTOR K mattress, super single", "FURNITURE", "MATTRESS", "PCS", ""),
    "CARRES L1201F MATT (Q)": ("CARRES L1201F mattress, queen", "FURNITURE", "MATTRESS", "PCS", ""),
    "CARRES B1201S MATT (Q)": ("CARRES B1201S mattress, queen", "FURNITURE", "MATTRESS", "PCS", ""),
    "FLOOR MATT (S)": ("Floor mattress, single", "FURNITURE", "MATTRESS", "PCS", ""),
    "FAN-FNZ-56-R-BK": ("FANZO ceiling fan 56 inch with remote, black", "LIGHTING", "CEILING_FAN", "UNIT", ""),
    "FAN-FNZ-42-R-BK": ("FANZO ceiling fan 42 inch with remote, black", "LIGHTING", "CEILING_FAN", "UNIT", ""),
    "FAN-W1L-DEKA-42-L": ("DEKA ceiling fan 42 inch with light", "LIGHTING", "CEILING_FAN", "UNIT", ""),
    "FAN-GEN-EXHAUST": ("Exhaust / ventilation fan", "LIGHTING", "CEILING_FAN", "UNIT", ""),
    "DLT-GEN-SURF": ("Surface downlight (size to confirm)", "LIGHTING", "DOWNLIGHT", "PCS", ""),
    "DLT-GEN-GRILL226": ("Grill light 226mm", "LIGHTING", "DOWNLIGHT", "PCS", ""),
    "DLT-GEN-GRILLADJ110": ("Adjustable grill light 110mm", "LIGHTING", "DOWNLIGHT", "PCS", ""),
    "PND-GEN-DESIGNER": ("Designer hanging lamp", "LIGHTING", "PENDANT", "PCS", ""),
    "TRK-GEN-1M-2P2": ("Track light 1m with 2+2 heads", "LIGHTING", "TRACK_LIGHT", "SET", ""),
    "TRK-GEN-2M-3": ("Track light 2m with 3 heads", "LIGHTING", "TRACK_LIGHT", "SET", ""),
    "TRK-GEN-LINEAR24W": ("LED linear track light 24W", "LIGHTING", "TRACK_LIGHT", "PCS", ""),
    "TRK-GEN-MAG-SURF": ("Surface magnetic track", "LIGHTING", "TRACK_LIGHT", "M", ""),
    "BLD-GEN-ZEBRA": ("Zebra blinds", "SOFT", "CURTAIN_BLINDS", "SQFT", ""),
    "BLD-GEN-ROLLER": ("Roller blinds", "SOFT", "CURTAIN_BLINDS", "SQFT", ""),
    "MIR-GEN-BRONZE": ("Bronze mirror", "SOFT", "DECORATION", "SQFT", ""),
    # services
    "SVC-CTS-DEFECT": ("Defect checking", "SERVICE", "", "LOT", ""),
    "SVC-CTS-DESIGN": ("Design fee / 2D and 3D drawing", "SERVICE", "", "LOT", ""),
    "SVC-CTS-SIGNAGE": ("Signboard / logo", "SERVICE", "", "LOT", ""),
}

# Owner, 26 Oct 2026: made-to-measure work, labour and unbranded goods use one code per sub-category (the line keeps its
# own words and price); a branded model or a package item that recurs (10+ lines) keeps its own code.
CATEGORY = {
    "RN-KITCHEN_CABINET": ("Built-in - Kitchen cabinet", "BUILTIN", "KITCHEN_CABINET", "FT", ""),
    "RN-WALL_CABINET": ("Built-in - Wall cabinet", "BUILTIN", "WALL_CABINET", "FT", ""),
    "RN-WARDROBE": ("Built-in - Wardrobe", "BUILTIN", "WARDROBE", "FT", ""),
    "RN-SHOE_CABINET": ("Built-in - Shoe cabinet", "BUILTIN", "SHOE_CABINET", "FT", ""),
    "RN-TV_CABINET": ("Built-in - TV cabinet", "BUILTIN", "TV_CABINET", "FT", ""),
    "RN-STUDY_TABLE": ("Built-in - Study table", "BUILTIN", "STUDY_TABLE", "FT", ""),
    "RN-SHELVING": ("Built-in - Shelving", "BUILTIN", "SHELVING", "FT", ""),
    "RN-COUNTERTOP": ("Built-in - Countertop (stone)", "BUILTIN", "COUNTERTOP", "FT", "stone colour is the LB SKU when known"),
    "RN-DESIGN": ("Service - Design", "SERVICE", "", "LOT", ""),
    "RN-SIGNAGE": ("Service - Signage", "SERVICE", "", "LOT", ""),
    "RN-DEFECT": ("Service - Defect checking", "SERVICE", "", "LOT", ""),
}
REMAP = {}
for old, new in [
    ("RN-KITCHEN_CABINET", "BIN-CTS-BASE BIN-CTS-TALL BIN-CTS-TALL-GLASS BIN-CTS-FRIDGE BIN-CTS-DRAWER-1 BIN-CTS-DRAWER-3 BIN-CTS-DRAWER-UM1 BIN-CTS-DRAWER-UM3 BIN-CTS-PULLOUT BIN-CTS-DISHRACK BIN-CTS-BACKPANEL BIN-CTS-FOAMBOARD BIN-CTS-EP"),
    ("RN-WALL_CABINET", "BIN-CTS-WALL"),
    ("RN-WARDROBE", "BIN-CTS-WARDROBE BIN-CTS-WARDROBE-3D BIN-CTS-HANGROD"),
    ("RN-SHOE_CABINET", "BIN-CTS-SHOE"),
    ("RN-TV_CABINET", "BIN-CTS-TV"),
    ("RN-STUDY_TABLE", "BIN-CTS-DRESSING BIN-CTS-STUDY"),
    ("RN-SHELVING", "BIN-CTS-SHELF BIN-CTS-BEDHEAD BIN-CTS-CABINET BIN-CTS-CURVE"),
    ("RN-COUNTERTOP", "BIN-CTS-QUARTZ BIN-CTS-SINTERED"),
    ("RN-ELECTRICAL", "ELC-CTS-LIGHT-PIPE ELC-CTS-LIGHT-CONCEAL ELC-CTS-SOCKET-1 ELC-CTS-SOCKET-2 ELC-CTS-SWITCH ELC-CTS-FANHOOK ELC-CTS-DOORBELL ELC-CTS-INST-FAN ELC-CTS-INST-LIGHT ELC-CTS-INST-HEATER ELC-CTS-INST-HOODHOB"),
    ("RN-DESIGN", "SVC-CTS-DESIGN"),
    ("RN-SIGNAGE", "SVC-CTS-SIGNAGE"),
    ("RN-DEFECT", "SVC-CTS-DEFECT"),

    ("RN-KITCHEN_CABINET", "BIN-CTS-ALUCASING"),
]:
    for o in new.split():
        assert o in NEW, o
        REMAP[o] = old
for o in REMAP:
    del NEW[o]
NEW.update(CATEGORY)
SIZES = [("K", "183X190"), ("Q", "152X190"), ("SS", "107X190"), ("S", "90X190")]
for old in [k for k in NEW if k.startswith("MAT-")]:
    del NEW[old]
for model, thick in [("AKEMI GALLENT", "20CM"), ("CARRES HOTEL PREMIUM L", ""), ("CARRES L1201F", ""), ("CARRES B1201S", ""),
                     ("DOCTOR K", ""), ("DR MACIO", "30CM"), ("FLOOR", "")]:
    for size, dims in SIZES:
        NEW[f"{model} MATT ({size})"] = (f"{model} MATTRESS ({dims}{'X' + thick if thick else 'CM'})", "FURNITURE", "MATTRESS", "PCS", "")
# Sofas as the catalog writes them: "SOF-HKA-5535-3S" / "SOFA 5535 3S"
for old, code, name in [("SOF-GEN-NEXON-3S", "SOF-GEN-NEXON-3S", "SOFA NEXON 3S"), ("SOF-GEN-BLATT-3S", "SOF-GEN-BLATT-3S", "SOFA BLATT 3S"),
                        ("SOF-GEN-PREM25", "SOF-GEN-PREMIUM-2.5S", "SOFA PREMIUM 2.5S"), ("SOF-GEN-3S-PUSHBACK", "SOF-GEN-PUSHBACK-3S", "SOFA PUSH BACK 3S"),
                        ("SOF-GEN-BED", "SOF-GEN-SOFABED", "SOFA BED"), ("SOF-GEN-1S", "SOF-GEN-1S", "SOFA 1S")]:
    v = NEW.pop(old)
    NEW[code] = (name,) + v[1:]
    REMAP[old] = code

SKIP = "SKIP"
REVIEW = "REVIEW"
R = [
    (SKIP, r"^(SUB ?)?TOTAL\b|^GRAND TOTAL|DISCOUNT|REBATE|^LESS\b|DEPOSIT|PAYMENT|BOOKING FEE"),
    # services
    ("RN-PROJECT_MGMT", r"PROJECT MANAGEMENT|PROJECT MGMT|SPACE PLANNING|SUPERVIS|MANAGEMENT FEE"),
    ("SVC-CTS-DEFECT", r"DEFECT"),
    ("RN-CLEANING", r"CLEANING|\bCLEAN\b|CLEAN UP"),
    ("RN-PHOTOSHOOT", r"PHOTO ?SHOOT"),
    ("RN-STYLING", r"STYLING|BORROW FOR SHOOTING"),
    ("RN-DELIVERY", r"^DELIVERY|TRANSPORT|LORRY"),
    ("SVC-CTS-SIGNAGE", r"SIGNBOARD|LOGO|LIGHT BOX|PRINTING"),
    ("SVC-CTS-DESIGN", r"DESIGN FEE|\b3D\b|DRAWING"),
    ("RN-AIRBNB_SETUP", r"AIRBNB|ABNB|HOUSEHOLD ITEM|STARTER KIT|^LINENS?\b|^AMENITIES"),
    ("LIN-0007", r"PILLOW ?CASE"),
    ("LIN-0006", r"SOFT PILLOW"),
    ("LIN-0004", r"BEDSHEET \(S"),
    ("LIN-0011", r"QUILT INSERT \(K"),
    ("LIN-0008", r"QUILT COVER \(K"),
    ("LIN-0013", r"SUPER SINGLE CON?FORTER"),
    ("AIB-0008", r"FLOOR RUGS?"),
    # installation of things the owner supplies
    ("ELC-CTS-INST-HOODHOB", r"(HOOD|HOB).*(BY OWNER|INSTALL)|INSTALL.*(HOOD|HOB)"),
    ("ELC-CTS-INST-HEATER", r"INSTALL.*HEATER|HEATER.*INSTALL"),
    ("RN-PLUMBING", r"INSTALL.*SINK"),
    ("ELC-CTS-INST-FAN", r"^(?!.*SUPPLY (&|AND) INSTALL)(?=.*INSTALL)(?=.*\bFANS?\b)"),
    ("ELC-CTS-INST-LIGHT", r"INSTALL ?(ONLY)? ?(CEILING )?LIGHT|INSTALL ONLY (SURFACE )?DOWN ?LIGHT|INSTALL ONLY LIGHTING|INSTALL.*LINEAR PROFILE"),
    # electrical
    ("ELC-CTS-FANHOOK", r"FAN HOOK"),
    ("ELC-CTS-DOORBELL", r"DOOR ?BELL"),
    ("ELC-CTS-SOCKET-2", r"SUBSEQ\w* SOCKET"),
    ("ELC-CTS-SOCKET-1", r"SOCKET|POWER ?POINT|13AMP"),
    ("ELC-CTS-LIGHT-CONCEAL", r"(LIGHT|LED) ?(ING)? ?POINT.*(CONCEAL|HACKING|DB|SWITCH)|NEW LIGHTPOINT"),
    ("ELC-CTS-LIGHT-PIPE", r"LIGHT(ING)? POINT"),
    ("ELC-CTS-SWITCH", r"\b1[- ]?GANG SWITCH"),
    ("RN-ELECTRICAL", r"TV BRACKET|SWITCH|GANG|WIRING|ELECTRIC|\bDB\b|CABLE|TRUNKING|ISOLATOR|DATA POINT|TV POINT|CCTV|M ?& ?E"),
    # appliances
    ("TV-HIS-50", r"(TELEVISION|\bTV\b).*50|50.*(SMART TV|TELEVISION)"),
    ("TV-HIS-43", r"(TELEVISION|\bTV\b).*43"),
    ("FRG-HIS-420L", r"FRIDGE 4[02]0 ?L"),
    ("FRG-GEN-110L", r"REFRIGERATOR 110"),
    ("FRG-GEN-200L", r"REFRIGERATOR 200"),
    ("FRG-GEN-280L", r"REFRIGERATOR 280"),
    ("CON-0005", r"DISHWASHER PASTE"),
    ("WSH-MID-2IN1", r"WASHING MACHINE|WASHER"),
    ("ACU-HIS-15HP", r"HISENSE 1\.5 ?HP|AIRCOND? HISENSE 1\.5"),
    ("ACU-HIS-1HP", r"HISENSE 1 ?HP|AIRCOND? HISENSE 1 HP"),
    ("RN-AIRCON_INSTALL", r"H/?P AIRCOND|AIRCON|AIR CON|COPPER"),
    ("HOB-HET-2IN1", r"HETCH|INDUCTION|COOKER"),
    ("HOD-ELB-M7598", r"ELBA|SMOKE HOOD"),
    ("HTR-GEN-HEATER", r"WATER HEATER"),
    # lighting
    ("FAN-GEN-5B56-RL", r"FAN 5 BLADE.*56.*LIGHT|CEILING FAN W/? ?LED"),
    ("FAN-GEN-5B56-R", r"FAN 5 BLADE.*56"),
    ("FAN-GEN-5B42-R", r"FAN 5 BLADE.*42"),
    ("FAN-GEN-WALL", r"WALL FAN"),
    ("FAN-FNZ-FANLIGHT", r"FANZO.*LIGHT"),
    ("FAN-FNZ-FAN", r"FANZO"),
    ("FAN-GEN-CEILING", r"CEILING FAN|LIGHTING / FANS"),
    ("ELC-CTS-LIGHT-PIPE", r"LOOPING NEW (LIGHT/ ?FAN )?POINT"),
    (REVIEW, r"\bFANS?\b"),
    ("DLT-GEN-SURF8-CW", r"SURFACE DOWN ?LI?GHT 8"),
    ("DLT-GEN-RECESS", r"RECESSED DOWN ?LI?GHT"),
    ("DLT-GEN-LINEAR600", r"LINEAR LIGHT"),
    ("DLT-GEN-SURF8-CW", r"SURFACE DOWN ?LI?GHT|DOWN ?LI?GHTS?\b"),
    ("LED-KNK-005", r"LED COB|COB STRIP|\bCOB\b"),
    ("BIN-CTS-ALUCASING", r"(ALU|ALUMINIUM)\w* (LED )?CASING|LED CASING|LED ALUMINIUM CASING"),
    ("TRK-GEN-MAGNETIC", r"MAGNETIC TRACK|MAGENETIC TRACK"),
    ("TRK-LTW-ZS-QL-POWER-TRACK-1000MM-WH", r"POWER TRACK"),
    ("TRK-GEN-1M-SET", r"TRACK LIGHT"),
    ("LGT-GEN-T5", r"\bT5\b"),
    ("EYB-GEN-EYEBALL", r"EYEBALL"),
    ("PND-GEN-DINING", r"HANGING (LIGHT|LAMP)|PENDANT"),
    ("PND-GEN-WALL", r"WALL (LIGHT|LAMP)"),
    # built-in carpentry, before the material words they contain (glass, laminate, door)
    ("BIN-CTS-TALL-GLASS", r"TALL ?CABINET.*GLASS"),
    ("BIN-CTS-FRIDGE", r"FRIDGE (CABINET|UNIT)|^FRIDGE \(|CABINET ABOVE FRIDGE"),
    ("BIN-CTS-TALL", r"TALL ?CABINET|OVEN CABINET|PANTRY"),
    ("BIN-CTS-DRAWER-UM3", r"UNDERMOU\w*.*DRAWER.*3"),
    ("BIN-CTS-DRAWER-UM1", r"UNDERMOU\w*.*DRAWER"),
    ("BIN-CTS-DRAWER-3", r"DRAWERS? *\(?3|3 LAYER"),
    ("BIN-CTS-DRESSING", r"^(BUILD- ?IN )?DRESSING TABLE"),
    ("BIN-CTS-WARDROBE", r"WARDROBE.*DRESSING"),
    ("BIN-CTS-DRAWER-1", r"DRAWER"),
    ("BIN-CTS-DISHRACK", r"DISH RACK"),
    ("BIN-CTS-PULLOUT", r"PULL[- ]?OUT|BASKET|LIFT UP|SAUCE TRAY"),
    ("FUR-MUJ-WARDROBE3D", r"MUJI 3 ?DOOR"),
    ("FUR-MUJ-WARDROBEOPEN", r"MUJI OPEN SHELF"),
    ("FUR-MUJ-TVCAB", r"MUJI.*TV"),
    ("FUR-MUJ-SIDETABLE", r"MUJI SIDE"),
    ("FUR-MUJ-COFFEESET", r"MUJI COFFEE"),
    ("FUR-MUJ-DINTABLE", r"MUJI DINING TABLE"),
    ("FUR-MUJ-DINCHAIR", r"MUJI DINING CHAIR"),
    ("FUR-MUJ-DINBENCH", r"MUJI DINING BENCH"),
    ("CUR-MUJ-DOUBLE", r"CURTAIN DOUBLE LAYER MUJI|MUJI.*DOUBLE"),
    ("CUR-MUJ-SINGLE", r"CURTAIN SINGLE LAYER MUJI|MUJI.*SINGLE"),
    (REVIEW, r"MUJI"),
    ("BIN-CTS-WARDROBE-3D", r"3 DOOR CLOSE WARDROBE"),
    ("BIN-CTS-WARDROBE", r"WARDRO?\w*|WARDORBE"),
    ("BIN-CTS-BASE", r"^ISLAND$|BASE (CABINET|UNIT)|BOTTOM CABINET|LOWER CABINET"),
    ("BIN-CTS-WALL", r"WALL (CABINET|UNIT)|TOP CABINET|UPPER CABINET|HANGING CABINET"),
    ("BIN-CTS-SHOE", r"SHOES? CABINET|ROTATING SHOE"),
    ("FIX-0004", r"^SHOE RACK"),
    ("FUR-GEN-DESKTV", r"WORKING DESK \+ TV"),
    ("BIN-CTS-TV", r"TV (CONSOLE|CABINET)|^CONSOLE|SITTING TV"),
    ("FUR-GEN-DRESSING", r"DESIGNER DRESSING"),
    ("BIN-CTS-DRESSING", r"DRESSING|VANITY"),
    ("FUR-GEN-WORKCHAIR", r"(WORKING|STUDY) CHAIR"),
    ("BIN-CTS-STUDY", r"STUDY (TABLE|DESK)|WORKING DESK|STUDY DESK|\bDESK\b"),
    ("BIN-CTS-HANGROD", r"HANGING ROD"),
    ("AIB-0010", r"CLOTH HANGING (RACK|PACK)"),
    ("BIN-CTS-SHELF", r"SHELF|SHELVES|SHELVING|BOOK ?SHELF|DISPLAY"),
    ("BIN-CTS-BEDHEAD", r"BED ?HEAD|HEADBOARD"),
    ("BIN-CTS-BACKPANEL", r"BACK(ING)? PANEL"),
    ("BIN-CTS-FOAMBOARD", r"FOAM BOARD"),
    ("BIN-CTS-EP", r"^MELAMINE EP|END PANEL"),
    ("BIN-CTS-CURVE", r"CU?RU?VE DESIGN"),
    ("BIN-CTS-SINTERED", r"SINTERED"),
    ("BIN-CTS-QUARTZ", r"QUARTZ|STONE (TOP|TABLE TOP)|KITCHEN TOP|SOLID SURFACE|BACK ?SPLASH"),
    ("RN-FEATURE_WALL", r"PU STONE PANEL"),
    ("BIN-CTS-CABINET", r"CABINET|STORAGE|CARPENTRY|BUILT? ?-?IN|PLYWOOD LAMINATE|LAMINATE PLYWOOD COUNTER|GAS SPRING|HANDLE|KICK|PLINTH|FILLER|HINGE|GAS STRUT"),
    # works (generic works SKUs keep the line's own text)
    ("RN-PROTECTION", r"PROTECTION|FLOOR GUARD"),
    ("RN-PAINTING", r"PAINT"),
    ("RN-FEATURE_WALL", r"FEATURE WALL|FLUT+ED|WAINSCOT|WALLPAPER|WALL PANEL"),
    ("RN-PLASTER_CEILING", r"PLASTER|CEILING|CORNICE|L ?BOX|BOX UP"),
    ("RN-HACKING", r"HACK|DEMOLISH|DISMANTLE|REMOVE|DISPOSAL|RUBBISH|DEBRIS"),
    ("RN-PLUMBING", r"PLUMB|PLUMPING|PIPE|PIPING|SINK|\bTAP\b|WATER POINT|BASIN|TOILET BOWL|\bWC\b|SHOWER|MIXER|FLOOR TRAP|BIDET"),
    ("RN-MASONRY_TILING", r"TILE|TILING|SCREED|BRICK|CEMENT|MASONRY|SKIRTING|KERB|CURB|SOFT POUR"),
    ("RN-WATERPROOFING", r"WATER ?PROOF"),
    ("RN-FLOORING", r"\bSPC\b|VINYL|LAMINATE|FLOORING|TIMBER FLOOR|PARQUET|TATAMI"),
    ("RN-GRILL", r"GRILL"),
    ("RN-PARTITION", r"PARTITION|GYPSUM|DRYWALL"),
    ("RN-ALUMINIUM_GLASS", r"GLASS|ALUMINIUM|WINDOW|SLIDING DOOR|SHOWER SCREEN|TEMPERED"),
    ("BLD-GEN-BLINDS", r"BLINDS?\b|ROLLER"),
    ("CUR-GEN-DOUBLE", r"DOUBLE[- ](PLY|LAYER)"),
    ("CUR-GEN-SINGLE", r"SINGL\w*[- ](PLY|LAYER)|CURTAIN|SHEER"),
    ("MIR-GEN-ROUND600", r"ROUND MIRROR|HANGING MIRROR"),
    ("MIR-CTS-CUSTOM", r"MIRROR"),
    ("RN-DOOR", r"\bDOOR\b|DIGITAL LOCK|SMART ?LOCK|DOOR LOCK|DOOR CLOSER"),
    # loose furniture and beds
    ("SOF-HKA-5535-L(LHF)", r"5535"),
    ("SOF-GEN-PREM25", r"PREMIUM SOFA 2\.5|2\.5 SEATER SOFA"),
    ("SOF-GEN-3S", r"3 ?SEATER.*SOFA|SOFA.*3 ?SEATER"),
    (REVIEW, r"SOFA"),
    ("AKEMI TRANQUIL MATT (K)", r"TRANQUIL"),
    ("AKEMI GALLENT MATT (Q)", r"GALLE?A?NT"),
    ("CARRES HOTEL PREMIUM L MATT (S)", r"CARRES HOTEL PREMIUM MODEL L"),
    (REVIEW, r"MATTRESS"),
    ("BED-HKA-1013-(Q)", r"JAGER.*(\(Q\)|QUEEN)|^JAGER BEDFRAME$"),
    ("BED-HKA-1013-(S)", r"JAGER.*\(S\)"),
    ("BED-HKA-1013-(SS)", r"JAGER.*SUPER SINGLE"),
    ("BED-HKA-1007-(K)", r"CODY.*KING|KING.*CODY"),
    ("BED-HKA-1007-(Q)", r"CODY"),
    (REVIEW, r"BED ?FRAME|DIVAN"),
    ("FUR-GEN-NEATH150", r"NEATH"),
    ("FUR-GEN-DINSET6", r"6 SEATER|1 ?\+ ?6"),
    ("FUR-GEN-DINSET2", r"DINING TABLE \+ 2"),
    (REVIEW, r"DINING TABLE"),
    ("FUR-GEN-CHAIR", r"CHAIR"),
    ("FUR-GEN-POUF", r"POUF|OTTOMAN"),
    ("FUR-GEN-COFFEE", r"COFFEE TABLE"),
    ("FUR-GEN-SIDETABLE", r"SIDE TABLE"),
    ("AIB-0008", r"CARPET|\bRUG\b"),
    ("DEC-0002", r"DECORATION ITEM|DECO ITEM"),
    ("DEC-0005", r"ART FRAME"),
]
# Checked before the balance file, whose short names ("BOWL", "SHOE RACK") would otherwise catch these.
FIRST = [
    ("RN-PLUMBING", r"TOILET BOWL"),
    ("FIX-0006", r"TOWEL (BAR )?RACK"),
    ("AIB-0029", r"MOP WITH BUCKET"),
    ("BIN-CTS-SHOE", r"ROTATING SHOE"),
    ("AKEMI GALLENT MATT (K)", r"GALLE?A?NT.*\(K\)|GALLE?A?NT.*KING"),
    ("AKEMI GALLENT MATT (Q)", r"GALLE?A?NT"),
    ("ACU-HIS-1HP", r"HISENSE AIR ?COND? 1 ?HP"),
    ("ELC-CTS-INST-HOODHOB", r"INSTALL.*(INDUCTION|COOKER)"),
    ("FUR-MUJ-TVCAB", r"MALMO"),
    ("BIN-CTS-FOAMBOARD", r"FOAM ?BOARD"),
    ("BIN-CTS-SHELF", r"WALL SHELF"),
    ("BED-HKA-1007-(SS)", r"CODY.*SUPER SINGLE"),
    ("TRK-LTW-IMI-TR16-12W-WH-4K", r"TRACK LIGHT 12W"),
    (REVIEW, r"STACKING KIT"),
    ("ELC-CTS-LIGHT-CONCEAL", r"LIGHT ?POINT.*(SOCKET|SWITCH)"),
    # install only is labour: the electrical works code
    ("ELC-CTS-INST-LIGHT", r"^(INSTALL|INSTALLATION)\b(?!.*SUPPLY).*(LIGHT|LAMP|DOWN ?LI?GHT|EYEBALL|TRACK|T5)|(LIGHT|EYEBALL|DOWNLIGHT|TRACK)\b.*(INSTALL ONLY|INSTALLATION ONLY|INSTALL \u2014)"),
    ("FUR-GEN-CHAIR-ZOLA", r"ZOLA"), ("FUR-GEN-CHAIR-NANZO", r"NANZO"), ("FUR-GEN-CHAIR-NAIDA", r"NAIDA"),
    ("FUR-GEN-CHAIR-MARLO", r"MARLO"), ("FUR-GEN-CHAIR-MACY", r"MACY"), ("FUR-GEN-CHAIR-CARTER", r"CARTER"),
    ("FUR-GEN-DIN-GERALD180", r"GERALD"), ("FUR-GEN-DIN-AIMON", r"AIMON"), ("FUR-GEN-DIN-ROUND180", r"ROUND DINING TABLE"),
    ("FUR-GEN-COF-SANDRA2", r"SANDRA"), ("FUR-GEN-COF-PANRA2", r"PANRA"), ("FUR-GEN-COF-OZZY80", r"OZZY"), ("FUR-GEN-COF-HIACE", r"HIACE"),
    ("FUR-GEN-SIDE-HIRADO", r"HIRADO"), ("FUR-GEN-BEDSIDE", r"BEDSIDE"),
    ("SOF-GEN-NEXON-3S", r"NEXON"), ("SOF-GEN-BLATT-3S", r"BLATT"), ("SOF-HKA-6112-3S", r"6112"),
    ("SOF-GEN-3S-PUSHBACK", r"PUSH BACK"), ("SOF-GEN-BED", r"SOFA BED"), ("SOF-GEN-1S", r"^SOFA 1S"),
    ("SOF-GEN-PREM25", r"SOFA 2\.5"),
    ("AKEMI OMNIA MATT (K)", r"OMNIA.*KING"), ("AKEMI OMNIA MATT (Q)", r"OMNIA.*QUEEN"),
    ("DR MACIO MATT (S)", r"MACIO"), ("DOCTOR K MATT (SS)", r"DOCTOR K.*\(\s*S\.?S"), ("DOCTOR K MATT (Q)", r"DOCTOR K"),
    ("CARRES L1201F MATT (Q)", r"L1201F"), ("CARRES B1201S MATT (Q)", r"B1201S"), ("FLOOR MATT (S)", r"FLOOR MATTRESS"),
    ("BED-HKA-DIVAN-(K)", r"DIVAN KING"), ("BED-HKA-DIVAN-(Q)", r"QUEEN DIVAN"),
    ("FAN-FNZ-56-R-BK", r"FANZO.*56"), ("FAN-FNZ-42-R-BK", r"FANZO.*42"), ("FAN-W1L-DEKA-42-L", r"DEKA.*LIGHT.*42"),
    ("FAN-GEN-EXHAUST", r"EXHAUST FAN|VENTILATION FAN"),
    ("DLT-GEN-GRILLADJ110", r"ADJ GRILL LIGHT"), ("DLT-GEN-GRILL226", r"GRILL LIGHT"),
    ("PND-GEN-DESIGNER", r"DESIGNER HANGING"),
    ("TRK-GEN-1M-2P2", r"TRACK LIGHT 1 ?M\w* 2 ?\+ ?2"), ("TRK-GEN-2M-3", r"2 ?M\w* TRACK ?LIGHT|TRACK ?LIGHT \(?2 ?M"),
    ("TRK-GEN-LINEAR24W", r"LINEAR TRACK LIGHT"), ("TRK-GEN-MAG-SURF", r"SURFACE MAGNETIC"),
    ("BLD-GEN-ZEBRA", r"ZEBRA"), ("BLD-GEN-ROLLER", r"ROLLER BLIND"), ("MIR-GEN-BRONZE", r"BRONZE MIRROR"),
    ("DLT-GEN-SURF8-CW", r"SURFACE DOWN ?LI?GHT 8"),
    ("DLT-GEN-SURF", r"SURFACE DOWN ?LI?GHT|^DOWN ?LIGHTS?$|WHOLE HOUSE LED DOWNLIGHT"),
]
FIRST = [(t, re.compile(rx)) for t, rx in FIRST]
RULES = [(t, re.compile(rx)) for t, rx in R]
for t, _ in R + [(t, None) for t, _ in FIRST]:
    assert t in (SKIP, REVIEW) or t in NEW or t in REMAP or t in cat or t in bal, t

def norm(s):
    return re.sub(r"\s+", " ", (s or "").strip().upper())

def core(s):
    s = re.sub(r"^[-*\u2022\s]+", "", s)
    return re.sub(r"\s*@.*$", "", s).strip()

def words(s):
    return [w.rstrip("S") for w in re.findall(r"[A-Z]{3,}", s)]

BAL_WORDS = {k: set(words(v[0])) for k, v in bal.items()}

def balance_match(c):
    """A balance SKU whose every word is in the line, the line having at most 2 words more."""
    have = set(words(c))
    best = None
    for k, bw in BAL_WORDS.items():
        if bw and bw <= have and len(have - bw) <= 2 and (best is None or len(bw) > len(BAL_WORDS[best])):
            best = k
    return best

def classify(k):
    c = core(k)
    if not c or re.fullmatch(r"[A-Z ]{0,20}(ROOM|KITCHEN|LIVING|DINING|BATHROOM|TOILET|BALCONY|YARD|ENTRANCE|FOYER|STUDY|AREA|HALL)S?\s*\d*:?", c):
        return SKIP
    if re.match(r"^(W|H|D|L)\s*:|^\(|^(DOUBLE|SINGLE) LAYER \d|^[A-Z' ]*(ROOM|LIVING|BEDROOM\d?|AREA)\s*-\s*[\d.]+\s*FT|^BACKSPLASH - |^OVAL SHAPE|^INSTALLATION OF$|^[A-Z0-9' ]*(ROOM|AREA|BEDROOM \d)\s+[\d.]+\s*FT$|^EXCLUDE ", c):
        return SKIP
    if RULES[0][1].search(c):
        return SKIP
    for t, rx in FIRST:
        if rx.search(c):
            return t
    b = balance_match(c)
    if b:
        return b
    for t, rx in RULES[1:]:
        if rx.search(c):
            return t
    return REVIEW

def kind(t):
    if t in NEW:
        return "New"
    if t in bal:
        return "Balance file"
    if t.startswith("RN-"):
        return "Existing (general)"
    return "Existing (exact)"

def label(t):
    if t in NEW:
        return NEW[t][0]
    if t in bal:
        return bal[t][0]
    return cat[t][1]

